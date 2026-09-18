"""Shared fixtures and helpers."""

from __future__ import annotations

from pathlib import Path
from typing import TYPE_CHECKING

import pytest

from resume_tailor.documents import parse

if TYPE_CHECKING:
    from docx.document import Document as DocxDocument

    from resume_tailor.documents.blocks import Document

REPO_ROOT = Path(__file__).resolve().parents[1]

RESUME_MD = """\
# Ada Lovelace
Principal Engineer
ada@example.com | (555) 010-0100 | London, UK | github.com/ada

## Summary
Engineer who writes **programs** for engines that do not exist yet.

## Skills
**Languages:** Analytical Notation, Mathematics
**Cloud & Infra:** AWS (Lambda, RDS), Kubernetes

## Experience

### Analytical Engine Programme — Principal Engineer
London, UK | Jan 1843 – Present
- **Algorithm Design:** Published the first algorithm intended for a machine.
- Corresponded with Babbage on engine semantics.
*Tech Stack — Bernoulli numbers, punched cards*

## Education

### Private tuition in mathematics
De Morgan | 1840

## Certifications
- Fellow of the Analytical Society — 1843
"""

LETTER_MD = """\
# Ada Lovelace
ada@example.com | London, UK

Dear Hiring Manager,

I would like to apply for the *engine* role.

Sincerely,
Ada
"""


@pytest.fixture
def resume() -> Document:
    """Return a parsed resume covering every block kind."""
    return parse(RESUME_MD)


@pytest.fixture
def letter() -> Document:
    """Return a parsed cover letter."""
    return parse(LETTER_MD)


def pdf_bytes(text: str) -> bytes:
    """Build a valid one-page PDF whose text layer is ``text`` (empty text stands in for a scan).

    Written by hand rather than with a PDF library: the point of the fixture is a *real* text
    layer to extract, and no writing library ships with this project.
    """
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1") if text else b""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream),
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n%s\nendobj\n" % (number, body)
    start = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        start,
    )
    return bytes(out)


def docx_lines(document: DocxDocument) -> list[str]:
    """Every non-empty paragraph in a rendered ``.docx``, including inside tables."""
    lines = [p.text.strip() for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                lines += [p.text.strip() for p in cell.paragraphs]
    return [line for line in lines if line]
