"""The ATS-safe layout — the properties a resume parser depends on."""

from __future__ import annotations

from typing import TYPE_CHECKING

from docx import Document as read_docx
from docx.shared import Pt

from resume_tailor.documents import parse
from resume_tailor.documents.blocks import Span, is_contact_line, is_note
from resume_tailor.render import ats
from tests.conftest import docx_lines

if TYPE_CHECKING:
    from pathlib import Path

    from docx.document import Document as DocxDocument
    from docx.text.paragraph import Paragraph as DocxParagraph

    from resume_tailor.documents.blocks import Document

# The entry-heading convention: the Markdown bolds the job title (or degree), never the company.
TITLED_MD = """\
# Ada Lovelace

## Experience
### **Principal Engineer** – Analytical Engine Programme
London, UK | Jan 1843 – Present

## Education
### **Mathematics** – Private tuition
1840
"""


def _render(document: Document, tmp_path: Path) -> DocxDocument:
    out = tmp_path / "resume.docx"
    ats.render_docx(document, out)
    return read_docx(str(out))


def _size(paragraph: DocxParagraph) -> float:
    size = paragraph.runs[0].font.size
    assert size is not None
    return float(size.pt)


def _style(paragraph: DocxParagraph) -> str:
    assert paragraph.style is not None
    return str(paragraph.style.name)


def test_docx_has_no_tables(resume: Document, tmp_path: Path) -> None:
    """Tables are the single most common cause of a scrambled parse."""
    assert not _render(resume, tmp_path).tables


def test_docx_keeps_contact_details_out_of_the_header_region(
    resume: Document, tmp_path: Path
) -> None:
    """Many parsers ignore the header/footer entirely."""
    document = _render(resume, tmp_path)
    for section in document.sections:
        assert not [p.text.strip() for p in section.header.paragraphs if p.text.strip()]
        assert not [p.text.strip() for p in section.footer.paragraphs if p.text.strip()]


def test_docx_contains_every_line_in_source_order(resume: Document, tmp_path: Path) -> None:
    lines = docx_lines(_render(resume, tmp_path))
    assert lines[0] == "Ada Lovelace"
    assert "ada@example.com | (555) 010-0100 | London, UK | github.com/ada" in lines
    assert "Summary" in lines
    assert "Languages: Analytical Notation, Mathematics" in lines
    assert "Analytical Engine Programme — Principal Engineer" in lines
    assert lines.index("Summary") < lines.index("Experience")


def test_section_headings_keep_their_case_as_the_design_does(
    resume: Document, tmp_path: Path
) -> None:
    lines = docx_lines(_render(resume, tmp_path))
    assert {"Summary", "Skills", "Experience", "Education", "Certifications"} <= set(lines)


def test_bullets_use_the_list_bullet_style(resume: Document, tmp_path: Path) -> None:
    document = _render(resume, tmp_path)
    matching = [p for p in document.paragraphs if "first algorithm" in p.text]
    styles = {p.style.name for p in matching if p.style is not None}
    assert styles == {"List Bullet"}


def test_inline_bold_becomes_a_bold_run(resume: Document, tmp_path: Path) -> None:
    document = _render(resume, tmp_path)
    runs = [r for p in document.paragraphs for r in p.runs if r.text == "programs"]
    assert runs and all(run.bold for run in runs)


def test_a_meta_line_renders_italic(resume: Document, tmp_path: Path) -> None:
    document = _render(resume, tmp_path)
    runs = [r for p in document.paragraphs for r in p.runs if "Jan 1843" in r.text]
    assert runs and all(run.italic for run in runs)


def test_every_run_uses_the_design_face(resume: Document, tmp_path: Path) -> None:
    document = _render(resume, tmp_path)
    fonts = {r.font.name for p in document.paragraphs for r in p.runs}
    assert fonts == {ats.FONT}


def test_the_header_sets_the_name_title_and_contact_line_apart(
    resume: Document, tmp_path: Path
) -> None:
    """A large regular name, a regular target title, a body-size regular contact line."""
    paragraphs = _render(resume, tmp_path).paragraphs
    name, title, contact = paragraphs[:3]
    assert (_size(name), name.runs[0].bold) == (28, False)
    assert name.paragraph_format.line_spacing == Pt(33)
    assert (_size(title), title.runs[0].bold) == (12, False)
    assert (_size(contact), contact.runs[0].bold) == (10, False)
    assert contact.text.startswith("ada@example.com |")


def test_a_heading_that_bolds_its_title_leaves_the_company_regular(tmp_path: Path) -> None:
    paragraphs = _render(parse(TITLED_MD), tmp_path).paragraphs
    role = next(p for p in paragraphs if p.text.startswith("Principal Engineer"))
    degree = next(p for p in paragraphs if p.text.startswith("Mathematics"))
    assert [(run.text, run.bold) for run in role.runs] == [
        ("Principal Engineer", True),
        (" – Analytical Engine Programme", False),
    ]
    assert [(run.text, run.bold) for run in degree.runs] == [
        ("Mathematics", True),
        (" – Private tuition", False),
    ]
    assert {run.font.name for p in (role, degree) for run in p.runs} == {ats.FONT}


def test_a_heading_with_no_bold_span_is_bold_throughout(resume: Document, tmp_path: Path) -> None:
    """Resumes written before the convention keep the all-bold heading they always had."""
    document = _render(resume, tmp_path)
    entry = next(p for p in document.paragraphs if "Analytical Engine Programme" in p.text)
    assert [(run.text, run.bold) for run in entry.runs] == [
        ("Analytical Engine Programme — Principal Engineer", True)
    ]


def test_bullets_hang_past_a_glyph_at_the_margin(resume: Document, tmp_path: Path) -> None:
    document = _render(resume, tmp_path)
    bullet = next(p for p in document.paragraphs if "first algorithm" in p.text)
    assert bullet.paragraph_format.left_indent == Pt(ats.BULLET_HANGING_IN * 72)
    assert bullet.paragraph_format.first_line_indent == Pt(-ats.BULLET_HANGING_IN * 72)


def test_the_first_bullet_of_a_run_stands_off_and_the_rest_sit_tight(
    resume: Document, tmp_path: Path
) -> None:
    document = _render(resume, tmp_path)
    first = next(p for p in document.paragraphs if "first algorithm" in p.text)
    second = next(p for p in document.paragraphs if "Corresponded" in p.text)
    assert first.paragraph_format.space_before == Pt(6.75)
    assert second.paragraph_format.space_before == Pt(2.25)


def test_a_closing_note_is_body_size_and_slightly_inset(resume: Document, tmp_path: Path) -> None:
    document = _render(resume, tmp_path)
    note = next(p for p in document.paragraphs if "Tech Stack" in p.text)
    assert _style(note) == "Normal"
    assert _size(note) == ats.BODY_PT
    assert all(run.italic for run in note.runs)
    assert note.paragraph_format.left_indent == Pt(ats.NOTE_INDENT_IN * 72)
    assert note.paragraph_format.space_before == Pt(2.25)


def test_prose_is_set_a_size_larger_than_the_bullets(resume: Document, tmp_path: Path) -> None:
    document = _render(resume, tmp_path)
    prose = next(p for p in document.paragraphs if "programs" in p.text)
    assert _size(prose) == ats.PROSE_PT
    assert prose.paragraph_format.line_spacing == Pt(ats.PROSE_LINE_PT)
    assert prose.paragraph_format.space_before == Pt(5.25)


def test_section_gaps_depend_on_what_precedes_them(resume: Document, tmp_path: Path) -> None:
    document = _render(resume, tmp_path)
    summary = next(p for p in document.paragraphs if p.text.strip() == "Summary")
    skills = next(p for p in document.paragraphs if p.text.strip() == "Skills")
    assert summary.paragraph_format.line_spacing == Pt(18)
    assert summary.paragraph_format.space_before == Pt(21), "follows the contact line"
    assert skills.paragraph_format.space_before == Pt(12), "follows the summary prose"

    bare = _render(parse("# Ada\n\n## Summary\nx"), tmp_path)
    assert bare.paragraphs[0].paragraph_format.space_before == Pt(0), "opens the page"
    assert bare.paragraphs[1].paragraph_format.space_before == Pt(37.5), "directly under the name"


def test_a_contact_line_is_any_header_line_with_a_pipe() -> None:
    assert is_contact_line((Span("a@b.c"), Span("|London")))
    assert not is_contact_line((Span("Principal Engineer"),))


def test_a_note_is_a_paragraph_set_entirely_in_italics() -> None:
    assert is_note((Span("Tech Stack — x", italic=True),))
    assert is_note((Span(" "), Span("x", italic=True)))
    assert not is_note((Span("Tech Stack — ", italic=True), Span("x")))
    assert not is_note(())


def test_a_cover_letter_renders_as_prose(letter: Document, tmp_path: Path) -> None:
    lines = docx_lines(_render(letter, tmp_path))
    assert lines[0] == "Ada Lovelace"
    assert "Dear Hiring Manager," in lines
    assert "Sincerely, Ada" in lines


def test_html_mirrors_the_docx_structure(resume: Document) -> None:
    html = ats.render_html(resume)
    assert html.startswith("<!doctype html>")
    assert "<title>Ada Lovelace</title>" in html
    assert "<h1>Ada Lovelace</h1>" in html
    assert "<h2>Summary</h2>" in html
    assert "<h3><strong>Analytical Engine Programme — Principal Engineer</strong></h3>" in html
    assert '<div class="meta"><em>London, UK | Jan 1843 – Present</em></div>' in html
    assert "<strong>Languages: </strong>Analytical Notation, Mathematics" in html


def test_html_wraps_consecutive_bullets_in_one_list(resume: Document) -> None:
    html = ats.render_html(resume)
    assert html.count("<ul>") == html.count("</ul>") == 2  # experience + certifications
    assert "<li><strong>Algorithm Design:</strong> Published" in html


def test_html_escapes_markup_in_content() -> None:
    html = ats.render_html(parse("# A <b>& Co</b>\n\n## Summary\n5 < 6"))
    assert "&lt;b&gt;" in html
    assert "5 &lt; 6" in html
    assert "<b>" not in html


def test_html_falls_back_to_a_generic_title() -> None:
    from resume_tailor.documents.blocks import Document as Blocks

    assert "<title>Resume</title>" in ats.render_html(Blocks())


def test_a_trailing_italic_note_renders_outside_the_bullet_list(resume: Document) -> None:
    html = ats.render_html(resume)
    assert '</ul><p class="note"><em>Tech Stack — Bernoulli numbers, punched cards</em></p>' in html


def test_html_carries_the_design_face_and_the_bullet_rule(resume: Document) -> None:
    html = ats.render_html(resume)
    assert "font-family: Arial, Helvetica, sans-serif;" in html
    assert "@import" not in html, "nothing to fetch: Arial is already installed"
    assert '<div class="subtitle">Principal Engineer</div>' in html
    assert '<div class="contact">ada@example.com |' in html
    assert f".contact {{ font-size: {ats.CONTACT_PT}pt; margin-top:" in html, "regular weight"
    assert 'li::before { content: "\u2022"; position: absolute; left: 0; }' in html
    assert "table" not in html.split("<body>")[1], "the single column never becomes a table"


def test_html_bolds_an_entry_heading_only_where_the_markdown_does() -> None:
    html = ats.render_html(parse(TITLED_MD))
    assert "<h3><strong>Principal Engineer</strong> – Analytical Engine Programme</h3>" in html
    assert "<h3><strong>Mathematics</strong> – Private tuition</h3>" in html
    assert f"h3 {{ font-size: {ats.BODY_PT}pt; font-weight: 400;" in html, "a regular base"
