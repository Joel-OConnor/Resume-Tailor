"""Orchestration: one Markdown file in, the right set of files out."""

from __future__ import annotations

import errno
import re
from typing import TYPE_CHECKING

import pytest

from resume_tailor.errors import RenderError
from resume_tailor.render import ats, exporter, polished
from resume_tailor.render.exporter import Layout, build
from resume_tailor.render.pdf import NO_BROWSER, PdfResult
from tests.conftest import RESUME_MD

if TYPE_CHECKING:
    from pathlib import Path

    from resume_tailor.documents.blocks import Document


@pytest.fixture
def source(tmp_path: Path) -> Path:
    path = tmp_path / "resume.md"
    path.write_text(RESUME_MD, encoding="utf-8")
    return path


@pytest.fixture(autouse=True)
def _pdf_always_succeeds(monkeypatch: pytest.MonkeyPatch) -> None:
    def fake(_: str, out: Path) -> PdfResult:
        out.write_bytes(b"%PDF-1.4\n")
        return PdfResult(ok=True)

    monkeypatch.setattr(exporter, "html_to_pdf", fake)


def _names(source: Path, **kwargs: object) -> list[str]:
    return [artifact.path.name for artifact in build(source, **kwargs)]  # type: ignore[arg-type]


def test_the_single_column_layout_by_default(source: Path) -> None:
    assert _names(source) == ["resume.docx", "resume.pdf"]


def test_both_layouts_on_request(source: Path) -> None:
    assert _names(source, layout=Layout.BOTH) == [
        "resume.docx",
        "resume.pdf",
        "resume-polished.docx",
        "resume-polished.pdf",
    ]


def test_the_ats_layout_keeps_the_source_stem(source: Path) -> None:
    assert _names(source, layout=Layout.ATS) == ["resume.docx", "resume.pdf"]


def test_the_polished_layout_is_suffixed(source: Path) -> None:
    assert _names(source, layout=Layout.POLISHED) == [
        "resume-polished.docx",
        "resume-polished.pdf",
    ]


def test_no_pdf_writes_only_word_files(source: Path) -> None:
    assert _names(source, pdf=False) == ["resume.docx"]
    assert _names(source, layout=Layout.BOTH, pdf=False) == ["resume.docx", "resume-polished.docx"]


def test_every_artifact_actually_exists(source: Path) -> None:
    assert all(artifact.path.is_file() for artifact in build(source))


def test_artifacts_report_their_layout(source: Path) -> None:
    assert [artifact.layout for artifact in build(source, layout=Layout.BOTH)] == [
        Layout.ATS,
        Layout.ATS,
        Layout.POLISHED,
        Layout.POLISHED,
    ]


def test_an_out_dir_is_created_and_used(source: Path, tmp_path: Path) -> None:
    out = tmp_path / "nested" / "exports"
    artifacts = build(source, out_dir=out, layout=Layout.ATS)
    assert all(artifact.path.parent == out for artifact in artifacts)
    assert (out / "resume.docx").is_file()


def test_a_failed_pdf_is_reported_as_html(source: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    def fails(html: str, out: Path) -> PdfResult:
        out.with_suffix(".html").write_text(html, encoding="utf-8")
        return PdfResult(ok=False, reason=NO_BROWSER)

    monkeypatch.setattr(exporter, "html_to_pdf", fails)
    artifacts = build(source, layout=Layout.ATS)
    assert [(a.path.name, a.ok) for a in artifacts] == [
        ("resume.docx", True),
        ("resume.html", False),
    ]


def test_the_sidebar_selection_reaches_the_polished_renderer(source: Path, tmp_path: Path) -> None:
    from docx import Document as read_docx

    build(source, layout=Layout.POLISHED, sidebar_sections=("summary",))
    table = read_docx(str(tmp_path / "resume-polished.docx")).tables[0]
    rail = [p.text.strip() for p in table.rows[0].cells[0].paragraphs if p.text.strip()]
    assert "Summary" in rail
    assert "Skills" not in rail


def test_each_pdf_is_printed_from_its_own_layouts_html(
    source: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """resume.pdf is the file people submit, so it must be the one column a parser can read."""
    printed: dict[str, str] = {}

    def record(html: str, out: Path) -> PdfResult:
        printed[out.name] = html
        out.write_bytes(b"%PDF-1.4\n")
        return PdfResult(ok=True)

    monkeypatch.setattr(exporter, "html_to_pdf", record)
    build(source, layout=Layout.BOTH, sidebar_sections=("summary",))
    document = exporter.read_source(source)

    assert set(printed) == {"resume.pdf", "resume-polished.pdf"}
    assert printed["resume.pdf"] == ats.render_html(document)
    assert "<aside" not in printed["resume.pdf"]
    # The polished PDF honours the rail choice, as its .docx twin does.
    assert printed["resume-polished.pdf"] == polished.render_html(document, ("summary",))
    rail = printed["resume-polished.pdf"].split('<aside class="rail">')[1].split("</aside>")[0]
    assert "<h2>Summary</h2>" in rail


def test_the_layout_enum_round_trips_through_strings() -> None:
    assert Layout("both") is Layout.BOTH
    assert str(Layout.ATS) == "ats"


def test_a_cover_letter_skips_the_polished_layout_even_when_both_are_asked_for(
    tmp_path: Path,
) -> None:
    """A letter has no sections: the rail would be empty, and the docs promise one column."""
    letter = tmp_path / "cover-letter.md"
    letter.write_text(
        "# Ada\nada@example.com\n\nDear Hiring Manager,\n\nHello.\n", encoding="utf-8"
    )
    assert _names(letter, layout=Layout.BOTH) == ["cover-letter.docx", "cover-letter.pdf"]


def test_an_explicit_polished_layout_still_renders_a_letter(tmp_path: Path) -> None:
    letter = tmp_path / "cover-letter.md"
    letter.write_text("# Ada\nada@example.com\n\nDear Hiring Manager,\n", encoding="utf-8")
    assert _names(letter, layout=Layout.POLISHED) == [
        "cover-letter-polished.docx",
        "cover-letter-polished.pdf",
    ]


def test_is_sectioned_distinguishes_a_resume_from_a_letter() -> None:
    from resume_tailor.documents import parse

    assert exporter.is_sectioned(parse("# A\n\n## Summary\nx"))
    assert not exporter.is_sectioned(parse("# A\n\nDear sir,\n"))


def test_a_dotted_stem_does_not_collapse_the_output_names(tmp_path: Path) -> None:
    """`with_suffix` would turn resume.v2.md into resume.docx, and polished would overwrite it."""
    source = tmp_path / "resume.v2.md"
    source.write_text(RESUME_MD, encoding="utf-8")
    names = _names(source, layout=Layout.BOTH, pdf=False)
    assert names == ["resume.v2.docx", "resume.v2-polished.docx"]
    assert (tmp_path / "resume.v2.docx").is_file()
    assert (tmp_path / "resume.v2-polished.docx").is_file()


def test_a_dotted_stem_keeps_the_ats_file_single_column(tmp_path: Path) -> None:
    from docx import Document as read_docx

    source = tmp_path / "jordan.rivera.md"
    source.write_text(RESUME_MD, encoding="utf-8")
    build(source, layout=Layout.BOTH, pdf=False)
    assert not read_docx(str(tmp_path / "jordan.rivera.docx")).tables
    assert read_docx(str(tmp_path / "jordan.rivera-polished.docx")).tables


def test_a_utf8_bom_does_not_hide_the_name_line(tmp_path: Path) -> None:
    source = tmp_path / "resume.md"
    source.write_bytes(b"\xef\xbb\xbf" + RESUME_MD.encode("utf-8"))
    assert _names(source, pdf=False) == ["resume.docx"]


def test_an_unreadable_source_is_a_render_error(tmp_path: Path) -> None:
    (tmp_path / "adir.md").mkdir()
    with pytest.raises(RenderError, match="cannot read"):
        build(tmp_path / "adir.md", pdf=False)


def test_a_non_utf8_source_is_a_render_error(tmp_path: Path) -> None:
    source = tmp_path / "resume.md"
    source.write_bytes(RESUME_MD.encode("utf-16"))
    with pytest.raises(RenderError, match="not UTF-8"):
        build(source, pdf=False)


def test_an_out_dir_that_is_a_file_is_a_render_error(source: Path, tmp_path: Path) -> None:
    blocker = tmp_path / "blocker"
    blocker.write_text("", encoding="utf-8")
    with pytest.raises(RenderError, match="cannot write to"):
        build(source, out_dir=blocker, pdf=False)


def test_an_unwritable_docx_target_is_a_render_error(source: Path, tmp_path: Path) -> None:
    (tmp_path / "resume.docx").mkdir()
    with pytest.raises(RenderError, match="cannot write"):
        build(source, pdf=False)


def test_a_pdf_write_failure_is_a_render_error(
    source: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def explode(_: str, __: Path) -> PdfResult:
        raise OSError(13, "Permission denied")

    monkeypatch.setattr(exporter, "html_to_pdf", explode)
    with pytest.raises(RenderError, match="cannot write"):
        build(source, layout=Layout.ATS)


def test_output_names_are_built_by_concatenation(tmp_path: Path) -> None:
    assert exporter.output_names(tmp_path, "resume.v2") == (
        tmp_path / "resume.v2.docx",
        tmp_path / "resume.v2.pdf",
        tmp_path / "resume.v2.html",
    )


def test_a_control_character_does_not_break_the_export(tmp_path: Path) -> None:
    source = tmp_path / "resume.md"
    source.write_text("# Ada\ny@z.com\n\n## Summary\nPasted\x0bfrom Word.\n", encoding="utf-8")
    assert _names(source, pdf=False) == ["resume.docx"]


def test_a_downgraded_letter_clears_the_previous_polished_export(tmp_path: Path) -> None:
    """Otherwise last build's two-column file waits in the folder the user sends from."""
    source = tmp_path / "resume.md"
    source.write_text(RESUME_MD, encoding="utf-8")
    build(source, layout=Layout.BOTH, pdf=False)
    assert (tmp_path / "resume-polished.docx").is_file()

    source.write_text("# Ada\ny@z.com\n\nDear Hiring Manager,\n", encoding="utf-8")
    assert _names(source, layout=Layout.BOTH, pdf=False) == ["resume.docx"]
    assert not (tmp_path / "resume-polished.docx").exists()


def test_an_explicit_layout_never_deletes_the_other_one(tmp_path: Path) -> None:
    """Only the tool's own downgrade cleans up; the user's choice is not destructive."""
    source = tmp_path / "resume.md"
    source.write_text(RESUME_MD, encoding="utf-8")
    build(source, layout=Layout.BOTH, pdf=False)
    build(source, layout=Layout.ATS, pdf=False)
    assert (tmp_path / "resume-polished.docx").is_file()


def test_no_pdf_clears_the_previous_pdf_and_its_fallback(tmp_path: Path) -> None:
    """Last build's resume.pdf would otherwise pass for the new resume.docx's twin."""
    source = tmp_path / "resume.md"
    source.write_text(RESUME_MD, encoding="utf-8")
    (tmp_path / "resume.pdf").write_bytes(b"%PDF-1.4\nLAST WEEK")
    (tmp_path / "resume.html").write_text("<p>last week</p>", encoding="utf-8")
    assert _names(source, pdf=False) == ["resume.docx"]
    assert sorted(path.name for path in tmp_path.iterdir()) == ["resume.docx", "resume.md"]


def test_no_pdf_leaves_a_layout_it_did_not_build_alone(tmp_path: Path) -> None:
    """Only the rebuilt layout's PDF is stale; an explicitly skipped layout keeps its files."""
    source = tmp_path / "resume.md"
    source.write_text(RESUME_MD, encoding="utf-8")
    build(source, layout=Layout.BOTH)
    build(source, layout=Layout.ATS, pdf=False)
    assert not (tmp_path / "resume.pdf").exists()
    assert (tmp_path / "resume-polished.docx").is_file()
    assert (tmp_path / "resume-polished.pdf").is_file()


def test_no_pdf_on_both_layouts_clears_both_previous_pdfs(tmp_path: Path) -> None:
    source = tmp_path / "resume.md"
    source.write_text(RESUME_MD, encoding="utf-8")
    build(source, layout=Layout.BOTH)
    assert _names(source, layout=Layout.BOTH, pdf=False) == ["resume.docx", "resume-polished.docx"]
    assert not (tmp_path / "resume.pdf").exists()
    assert not (tmp_path / "resume-polished.pdf").exists()


# --- a source is never one of its own outputs -----------------------------------------------------
@pytest.mark.parametrize(
    ("name", "pdf"),
    [
        ("resume.html", True),
        ("resume.html", False),
        ("resume.pdf", True),
        ("resume.pdf", False),
        ("resume.docx", False),
    ],
)
def test_a_source_saved_under_an_output_name_is_refused_untouched(
    tmp_path: Path, name: str, *, pdf: bool
) -> None:
    """Markdown saved as resume.html was deleted, and as resume.docx overwritten, with exit 0."""
    source = tmp_path / name
    source.write_text(RESUME_MD, encoding="utf-8")
    expected = rf"output of the build \({re.escape(name)}\).*rename it to resume\.md"
    with pytest.raises(RenderError, match=expected):
        build(source, pdf=pdf)
    assert source.read_text(encoding="utf-8") == RESUME_MD
    assert [path.name for path in tmp_path.iterdir()] == [name], "nothing is written or removed"


def test_both_layouts_check_the_polished_names_too(tmp_path: Path) -> None:
    source = tmp_path / "resume.html"
    source.write_text(RESUME_MD, encoding="utf-8")
    with pytest.raises(RenderError, match="destroy"):
        build(source, layout=Layout.BOTH, pdf=False)
    assert source.read_text(encoding="utf-8") == RESUME_MD


def test_a_layout_whose_names_cannot_clash_still_builds(tmp_path: Path) -> None:
    """The polished pass writes resume-polished.*, so a resume.html source is safe from it."""
    source = tmp_path / "resume.html"
    source.write_text(RESUME_MD, encoding="utf-8")
    assert _names(source, layout=Layout.POLISHED, pdf=False) == ["resume-polished.docx"]
    assert source.read_text(encoding="utf-8") == RESUME_MD


def test_a_source_differing_from_an_output_only_in_case_survives(tmp_path: Path) -> None:
    """On macOS and Windows resume.HTML and resume.html are one file, and removing one is both."""
    source = tmp_path / "resume.HTML"
    source.write_text(RESUME_MD, encoding="utf-8")
    if (tmp_path / "resume.html").exists():  # a disk that ignores case, as macOS's does
        with pytest.raises(RenderError, match="destroy"):
            build(source, pdf=False)
    else:  # a case-sensitive disk, where the two are separate files
        build(source, pdf=False)
    assert source.read_text(encoding="utf-8") == RESUME_MD


# --- a failed write never leaves a broken file ----------------------------------------------------
def _half_written(_: Document, out: Path) -> None:
    """Start the file and then fail, as a full disk does partway through a write."""
    out.write_bytes(b"PK\x03\x04 the first part of a .docx and nothing more")
    raise OSError(errno.ENOSPC, "No space left on device")


def test_a_failed_docx_write_keeps_the_previous_file_whole(
    source: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Writing in place, a full disk left a truncated resume.docx where the good one had been."""
    build(source, pdf=False)
    good = (tmp_path / "resume.docx").read_bytes()

    monkeypatch.setattr(ats, "render_docx", _half_written)
    with pytest.raises(RenderError, match=r"cannot write .*resume\.docx: No space left on device"):
        build(source, pdf=False)
    assert (tmp_path / "resume.docx").read_bytes() == good
    assert sorted(path.name for path in tmp_path.iterdir()) == ["resume.docx", "resume.md"]


def test_a_failed_first_docx_write_leaves_nothing_behind(
    source: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(ats, "render_docx", _half_written)
    with pytest.raises(RenderError, match="No space left on device"):
        build(source, pdf=False)
    assert [path.name for path in tmp_path.iterdir()] == ["resume.md"]


def test_the_docx_is_swapped_in_with_ordinary_permissions(source: Path, tmp_path: Path) -> None:
    """A temporary file is owner-only; the resume must be readable as any new file in the folder."""
    build(source, pdf=False)
    probe = tmp_path / "probe"
    probe.write_bytes(b"")
    assert (tmp_path / "resume.docx").stat().st_mode == probe.stat().st_mode
