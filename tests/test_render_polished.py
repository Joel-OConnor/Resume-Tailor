"""The two-column polished layout."""

from __future__ import annotations

from typing import TYPE_CHECKING

import pytest
from docx import Document as read_docx
from docx.shared import Inches, Pt

from resume_tailor.documents import parse
from resume_tailor.documents.blocks import Document as Blocks
from resume_tailor.documents.blocks import HeaderLine, SectionGroup, Span
from resume_tailor.render import polished

if TYPE_CHECKING:
    from pathlib import Path

    from docx.table import Table

    from resume_tailor.documents.blocks import Document


_TWIP = 635  # EMU

# The entry-heading convention: the Markdown bolds the job title (or degree), never the company.
# Education goes to the rail, so the rule is exercised in both cells.
TITLED_MD = """\
# Ada Lovelace

## Experience
### **Principal Engineer** – Analytical Engine Programme
London, UK | Jan 1843 – Present

## Education
### **Mathematics** – Private tuition
1840
"""


def _table(document: Document, tmp_path: Path, **kwargs: object) -> Table:
    out = tmp_path / "resume-polished.docx"
    polished.render_docx(document, out, **kwargs)  # type: ignore[arg-type]
    return read_docx(str(out)).tables[0]


def _cell_lines(table: Table, index: int) -> list[str]:
    return [p.text.strip() for p in table.rows[0].cells[index].paragraphs if p.text.strip()]


def _rail_spans(sidebar: list[SectionGroup]) -> list[tuple[Span, ...]]:
    """Return the preamble blocks the rail received, as spans."""
    return [block.spans for block in sidebar[0].blocks if isinstance(block, HeaderLine)]


def _rail_text(sidebar: list[SectionGroup]) -> list[str]:
    return [spans[0].text for spans in _rail_spans(sidebar)]


# --- column assignment ----------------------------------------------------------------------------
def test_skills_and_education_go_to_the_rail(resume: Document) -> None:
    sidebar, main = polished.split_columns(resume)
    assert [group.title for group in sidebar] == ["", "Skills", "Education", "Certifications"]
    assert [group.title for group in main] == ["", "Summary", "Experience"]


def test_sidebar_selection_is_configurable(resume: Document) -> None:
    sidebar, main = polished.split_columns(resume, ("summary",))
    assert [group.title for group in sidebar] == ["", "Summary"]
    assert "Skills" in [group.title for group in main]


def test_sidebar_matching_ignores_case(resume: Document) -> None:
    sidebar, _ = polished.split_columns(resume, ("SKILLS",))
    assert "Skills" in [group.title for group in sidebar]


def test_the_contact_line_moves_to_the_rail_one_item_per_line(resume: Document) -> None:
    sidebar, main = polished.split_columns(resume)
    assert _rail_text(sidebar) == [
        "ada@example.com",
        "(555) 010-0100",
        "London, UK",
        "github.com/ada",
    ]
    # The name and the target-title line stay in the main column.
    assert [type(block).__name__ for block in main[0].blocks] == ["Name", "HeaderLine"]


def test_empty_contact_items_are_dropped() -> None:
    document = Blocks((HeaderLine((Span("a || b |"),)),))
    sidebar, _ = polished.split_columns(document)
    assert _rail_text(sidebar) == ["a", "b"]


def test_contact_emphasis_survives_the_split() -> None:
    document = Blocks((HeaderLine((Span("a | ", bold=True), Span("b", italic=True))),))
    sidebar, _ = polished.split_columns(document)
    assert _rail_spans(sidebar) == [
        (Span("a", bold=True),),
        (Span("b", italic=True),),
    ]


# --- .docx ----------------------------------------------------------------------------------------
def test_docx_is_a_single_two_column_table(resume: Document, tmp_path: Path) -> None:
    table = _table(resume, tmp_path)
    assert len(table.rows) == 1
    assert len(table.columns) == 2
    # Widths round-trip through twips, so compare to the nearest twip rather than exactly.
    assert abs(table.columns[0].width - Inches(polished.SIDEBAR_WIDTH_IN)) <= _TWIP
    assert abs(table.columns[1].width - Inches(polished.MAIN_WIDTH_IN)) <= _TWIP


def test_docx_columns_span_the_full_page_width(resume: Document, tmp_path: Path) -> None:
    out = tmp_path / "resume-polished.docx"
    polished.render_docx(resume, out)
    document = read_docx(str(out))
    section = document.sections[0]
    assert section.page_width is not None
    assert section.left_margin is not None
    assert section.right_margin is not None
    usable = section.page_width - section.left_margin - section.right_margin
    assert usable == sum(column.width for column in read_docx(str(out)).tables[0].columns)


def test_docx_draws_the_divider_on_the_rail(resume: Document, tmp_path: Path) -> None:
    table = _table(resume, tmp_path)
    xml = table.rows[0].cells[0]._tc.xml
    assert "w:tcBorders" in xml
    assert 'w:right w:val="single" w:sz="8"' in xml
    # ...and only on the rail: a border on the main cell would double the rule.
    assert "w:tcBorders" not in table.rows[0].cells[1]._tc.xml


def test_docx_has_no_leading_blank_paragraph(resume: Document, tmp_path: Path) -> None:
    table = _table(resume, tmp_path)
    assert table.rows[0].cells[0].paragraphs[0].text.strip()
    assert table.rows[0].cells[1].paragraphs[0].text.strip() == "Ada Lovelace"


def test_rail_holds_contact_skills_and_education(resume: Document, tmp_path: Path) -> None:
    lines = _cell_lines(_table(resume, tmp_path), 0)
    assert lines[:4] == ["ada@example.com", "(555) 010-0100", "London, UK", "github.com/ada"]
    assert "Skills" in lines
    assert "Education" in lines
    assert "Analytical Engine Programme — Principal Engineer" not in lines


def test_main_holds_the_name_summary_and_experience(resume: Document, tmp_path: Path) -> None:
    lines = _cell_lines(_table(resume, tmp_path), 1)
    assert lines[0] == "Ada Lovelace"
    assert lines[1] == "Principal Engineer"
    assert "Summary" in lines
    assert "Analytical Engine Programme — Principal Engineer" in lines
    assert "Skills" not in lines


def test_rail_stacks_skill_items_one_per_line(resume: Document, tmp_path: Path) -> None:
    lines = _cell_lines(_table(resume, tmp_path), 0)
    assert "Languages" in lines
    assert "Analytical Notation" in lines
    assert "Mathematics" in lines


def test_commas_inside_parentheses_do_not_split_a_skill(resume: Document, tmp_path: Path) -> None:
    lines = _cell_lines(_table(resume, tmp_path), 0)
    assert "AWS (Lambda, RDS)" in lines
    assert "Kubernetes" in lines


def test_a_skill_line_in_the_main_column_stays_inline(resume: Document, tmp_path: Path) -> None:
    lines = _cell_lines(_table(resume, tmp_path, sidebar_sections=("summary",)), 1)
    assert "Languages: Analytical Notation, Mathematics" in lines


def test_skill_items_keep_their_emphasis() -> None:
    items = (Span("a, "), Span("b", bold=True), Span(", c"))
    assert polished._skill_items(items) == [
        (Span("a"),),
        (Span("b", bold=True),),
        (Span("c"),),
    ]


def test_unbalanced_parentheses_do_not_swallow_the_rest() -> None:
    assert polished._skill_items((Span("a), b"),)) == [(Span("a)"),), (Span("b"),)]


def test_the_name_is_large_and_regular_weight(resume: Document, tmp_path: Path) -> None:
    table = _table(resume, tmp_path)
    run = table.rows[0].cells[1].paragraphs[0].runs[0]
    assert run.font.name == polished.FONT
    assert run.font.size.pt == polished.NAME_PT
    assert run.bold is False


def test_a_heading_that_bolds_its_title_leaves_the_company_regular(tmp_path: Path) -> None:
    rail, main = _table(parse(TITLED_MD), tmp_path).rows[0].cells
    role = next(p for p in main.paragraphs if p.text.startswith("Principal Engineer"))
    degree = next(p for p in rail.paragraphs if p.text.startswith("Mathematics"))
    assert [(run.text, run.bold) for run in role.runs] == [
        ("Principal Engineer", True),
        (" – Analytical Engine Programme", False),
    ]
    assert [(run.text, run.bold) for run in degree.runs] == [
        ("Mathematics", True),
        (" – Private tuition", False),
    ]
    assert {run.font.name for p in (role, degree) for run in p.runs} == {polished.FONT}


def test_a_heading_with_no_bold_span_is_bold_throughout(resume: Document, tmp_path: Path) -> None:
    """Resumes written before the convention keep the all-bold heading they always had."""
    rail, main = _table(resume, tmp_path).rows[0].cells
    role = next(p for p in main.paragraphs if "Analytical Engine Programme" in p.text)
    tuition = next(p for p in rail.paragraphs if "Private tuition" in p.text)
    assert [(run.text, run.bold) for run in role.runs] == [
        ("Analytical Engine Programme — Principal Engineer", True)
    ]
    assert [(run.text, run.bold) for run in tuition.runs] == [
        ("Private tuition in mathematics", True)
    ]


def test_every_bullet_keeps_the_glyph_lead_in_or_not(resume: Document, tmp_path: Path) -> None:
    """The design bullets every accomplishment, bold lead-in included."""
    table = _table(resume, tmp_path)
    main = table.rows[0].cells[1].paragraphs
    lead = next(p for p in main if "first algorithm" in p.text)
    plain = next(p for p in main if "Corresponded" in p.text)
    assert lead.style.name == plain.style.name == "List Bullet"
    assert lead.paragraph_format.first_line_indent == Pt(-polished.BULLET_HANGING_IN * 72)
    assert lead.paragraph_format.left_indent == Pt(
        (polished.MAIN_INDENT[0] + polished.BULLET_HANGING_IN) * 72
    )


def test_the_first_bullet_of_a_run_stands_off_and_the_rest_sit_tight(
    resume: Document, tmp_path: Path
) -> None:
    table = _table(resume, tmp_path)
    main = table.rows[0].cells[1].paragraphs
    lead = next(p for p in main if "first algorithm" in p.text)
    plain = next(p for p in main if "Corresponded" in p.text)
    assert lead.paragraph_format.space_before == Pt(6.75)
    assert plain.paragraph_format.space_before == Pt(2.25)


def test_rail_skill_items_are_bulleted_with_the_rail_spacing(
    resume: Document, tmp_path: Path
) -> None:
    table = _table(resume, tmp_path)
    rail = table.rows[0].cells[0].paragraphs
    first = next(p for p in rail if p.text.strip() == "Analytical Notation")
    second = next(p for p in rail if p.text.strip() == "Mathematics")
    label = next(p for p in rail if p.text.strip() == "Languages")
    assert first.style.name == second.style.name == "List Bullet"
    assert first.paragraph_format.space_before == Pt(3.75)
    assert second.paragraph_format.space_before == Pt(1.5)
    assert label.style.name == "Normal"
    assert label.runs[0].bold is True


def test_a_closing_note_is_body_size_italic_and_slightly_inset(
    resume: Document, tmp_path: Path
) -> None:
    table = _table(resume, tmp_path)
    note = next(p for p in table.rows[0].cells[1].paragraphs if "Tech Stack" in p.text)
    assert note.runs[0].font.size.pt == polished.BODY_PT
    assert note.paragraph_format.space_before == Pt(2.25)
    assert note.paragraph_format.left_indent == Pt(
        (polished.MAIN_INDENT[0] + polished.NOTE_INDENT_IN) * 72
    )


def test_prose_is_set_a_size_larger_than_the_bullets(resume: Document, tmp_path: Path) -> None:
    table = _table(resume, tmp_path)
    prose = next(p for p in table.rows[0].cells[1].paragraphs if "programs" in p.text)
    assert prose.runs[0].font.size.pt == polished.PROSE_PT
    assert prose.paragraph_format.line_spacing == Pt(polished.PROSE_LINE_PT)
    assert prose.paragraph_format.space_before == Pt(5.25)


def test_each_column_opens_at_the_page_margin(resume: Document, tmp_path: Path) -> None:
    table = _table(resume, tmp_path)
    rail, main = table.rows[0].cells
    assert rail.paragraphs[0].paragraph_format.space_before == Pt(0)
    assert rail.paragraphs[1].paragraph_format.space_before == Pt(4.5)
    assert main.paragraphs[0].paragraph_format.space_before == Pt(0)


def test_section_gaps_depend_on_what_precedes_them(resume: Document, tmp_path: Path) -> None:
    table = _table(resume, tmp_path)
    rail, main = table.rows[0].cells
    skills = next(p for p in rail.paragraphs if p.text.strip() == "Skills")
    summary = next(p for p in main.paragraphs if p.text.strip() == "Summary")
    experience = next(p for p in main.paragraphs if p.text.strip() == "Experience")
    assert skills.paragraph_format.space_before == Pt(15.75)
    assert summary.paragraph_format.space_before == Pt(21.0), "follows the subtitle"
    assert experience.paragraph_format.space_before == Pt(9.75), "follows the summary prose"

    bare = parse("# Ada\n\n## Summary\nx")
    heading = _table(bare, tmp_path).rows[0].cells[1].paragraphs[1]
    assert heading.text.strip() == "Summary"
    assert heading.paragraph_format.space_before == Pt(37.5), "directly under the name"


def test_a_cover_letter_renders_entirely_in_the_main_column(
    letter: Document, tmp_path: Path
) -> None:
    table = _table(letter, tmp_path)
    assert "Dear Hiring Manager," in _cell_lines(table, 1)
    assert _cell_lines(table, 0) == ["ada@example.com", "London, UK"]


# --- HTML -----------------------------------------------------------------------------------------
def test_html_has_both_columns(resume: Document) -> None:
    html = polished.render_html(resume)
    assert '<aside class="rail">' in html
    assert '<section class="main">' in html
    assert f"width: {polished.SIDEBAR_WIDTH_IN}in" in html


def test_html_places_content_in_the_matching_column(resume: Document) -> None:
    html = polished.render_html(resume)
    rail = html.split('<aside class="rail">')[1].split("</aside>")[0]
    main = html.split('<section class="main">')[1].split("</section>")[0]
    assert '<div class="contact">ada@example.com</div>' in rail
    assert '<div class="skill-group">Languages</div>' in rail
    assert "<h2>Skills</h2>" in rail
    assert "<h1>Ada Lovelace</h1>" in main
    assert '<div class="subtitle">Principal Engineer</div>' in main
    assert "<h2>Experience</h2>" in main
    assert "<li><strong>Algorithm Design:</strong> Published" in main
    assert "<li>Analytical Notation</li>" in rail


def test_html_lists_open_and_close_in_pairs(resume: Document) -> None:
    html = polished.render_html(resume)
    # One list per run of bullets, one per rail skill group.
    assert html.count("<ul>") == html.count("</ul>") == 4
    assert "<li>Corresponded with Babbage on engine semantics.</li>" in html


def test_html_draws_every_bullet_glyph_at_the_column_edge(resume: Document) -> None:
    html = polished.render_html(resume)
    assert 'li::before { content: "\u2022"; position: absolute; left: 0; }' in html
    assert f"li {{ position: relative; padding-left: {polished.BULLET_HANGING_IN}in" in html


def test_html_sets_arial_and_fetches_nothing(resume: Document) -> None:
    html = polished.render_html(resume)
    assert "font-family: Arial, Helvetica, sans-serif;" in html
    assert "@import" not in html, "nothing to fetch: Arial is already installed"
    assert f".contact {{ font-size: {polished.CONTACT_PT}pt; margin-top:" in html, "regular weight"


def test_html_bolds_an_entry_heading_only_where_the_markdown_does() -> None:
    html = polished.render_html(parse(TITLED_MD))
    rail = html.split('<aside class="rail">')[1].split("</aside>")[0]
    main = html.split('<section class="main">')[1].split("</section>")[0]
    assert "<h3><strong>Principal Engineer</strong> – Analytical Engine Programme</h3>" in main
    assert "<h3><strong>Mathematics</strong> – Private tuition</h3>" in rail
    assert f"h3 {{ font-size: {polished.BODY_PT}pt; font-weight: 400;" in html, "a regular base"


def test_html_sets_a_heading_with_no_bold_span_bold_throughout(resume: Document) -> None:
    html = polished.render_html(resume)
    assert "<h3><strong>Analytical Engine Programme — Principal Engineer</strong></h3>" in html
    assert "<h3><strong>Private tuition in mathematics</strong></h3>" in html


def test_html_marks_a_closing_note_and_sets_prose_larger(resume: Document) -> None:
    html = polished.render_html(resume)
    assert '<p class="note"><em>Tech Stack' in html
    assert "<p>Engineer who writes" in html
    assert f"p {{ font-size: {polished.PROSE_PT}pt; line-height: {polished.PROSE_LINE_PT}pt" in html


def test_html_inline_skills_when_the_section_is_in_the_main_column(resume: Document) -> None:
    html = polished.render_html(resume, ("summary",))
    assert '<div class="skill-inline"><strong>Languages: </strong>' in html


def test_html_escapes_content() -> None:
    html = polished.render_html(parse("# A <b>\n\n## Summary\n5 < 6"))
    assert "&lt;b&gt;" in html
    assert "5 &lt; 6" in html


def test_html_falls_back_to_a_generic_title() -> None:
    assert "<title>Resume</title>" in polished.render_html(Blocks())


def test_the_rule_is_not_forced_to_full_page_height(resume: Document) -> None:
    """The .docx table stops at the content; the PDF has to agree."""
    assert "min-height" not in polished.render_html(resume)


@pytest.mark.parametrize("section", polished.DEFAULT_SIDEBAR_SECTIONS)
def test_every_default_rail_section_is_recognised(section: str) -> None:
    document = parse(f"# A\n\n## {section.title()}\n- x")
    sidebar, _ = polished.split_columns(document)
    assert [group.title for group in sidebar][1:] == [section.title()]


def test_empty_skill_items_are_dropped() -> None:
    assert polished._skill_items((Span("a,, b,"),)) == [(Span("a"),), (Span("b"),)]


def test_a_skill_line_with_no_items_renders_only_the_group_heading(tmp_path: Path) -> None:
    document = parse("# A\n\n## Skills\n**Languages:**")
    lines = _cell_lines(_table(document, tmp_path), 0)
    assert lines == ["Skills", "Languages"]


def test_a_trailing_italic_note_renders_without_a_bullet_glyph(
    resume: Document, tmp_path: Path
) -> None:
    """The `*Tech Stack — …*` line closes a role; a glyph would read as another accomplishment."""
    table = _table(resume, tmp_path)
    paragraphs = [p for p in table.rows[0].cells[1].paragraphs if "Tech Stack" in p.text]
    assert [p.style.name for p in paragraphs] == ["Normal"]
    assert all(run.italic for p in paragraphs for run in p.runs)


def test_consecutive_plain_bullets_share_one_html_list() -> None:
    html = polished.render_html(parse("# A\n\n## Awards\n- one\n- two\n- three"))
    assert html.count("<ul>") == html.count("</ul>") == 1
    assert html.count("<li>") == 3


def test_an_empty_rail_keeps_a_placeholder_paragraph(tmp_path: Path) -> None:
    """A `<w:tc>` with no block-level child is invalid per ECMA-376 and Word refuses to open it."""
    document = parse("# Ada\n\n## Summary\nno contact line, no rail sections")
    table = _table(document, tmp_path)
    rail = table.rows[0].cells[0]
    assert rail.paragraphs
    assert not rail.paragraphs[0].text.strip()
    assert "<w:p" in rail._tc.xml


def test_the_html_divider_matches_the_docx_border_width(resume: Document) -> None:
    """The .docx draws 1pt (w:sz=8 eighth-points); the PDF has to draw the same rule."""
    assert "border-right: 1pt solid #000" in polished.render_html(resume)


def test_any_pipe_marks_a_contact_line_as_the_contract_now_says() -> None:
    """Documented as "any pipe": a spaceless contact line must still reach the rail."""
    document = parse("# Ada\nPrincipal Engineer\na@b.c|London\n\n## Summary\nx")
    sidebar, main = polished.split_columns(document)
    assert _rail_text(sidebar) == ["a@b.c", "London"]
    assert [b.spans[0].text for b in main[0].blocks if isinstance(b, HeaderLine)] == [
        "Principal Engineer"
    ]


def test_a_list_that_ends_a_section_is_closed_before_the_next_heading() -> None:
    html = polished.render_html(parse("# A\n\n## Awards\n- one\n- two\n\n## Projects\n- three"))
    assert html.count("<ul>") == html.count("</ul>") == 2
    assert "<li>two</li></ul><h2>Projects</h2><ul><li>three</li></ul>" in html


def test_the_normal_style_carries_the_body_face(resume: Document, tmp_path: Path) -> None:
    """Word sizes a line by its paragraph mark too, and the mark takes Normal's font."""
    out = tmp_path / "resume-polished.docx"
    polished.render_docx(resume, out)
    normal = read_docx(str(out)).styles["Normal"].font
    assert normal.name == "Arial"
    assert normal.size == Pt(10)
