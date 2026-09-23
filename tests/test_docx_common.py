"""The low-level ``python-docx`` helpers, exercised directly."""

from __future__ import annotations

from typing import TYPE_CHECKING

from docx import Document as new_docx
from docx.oxml.ns import qn
from docx.shared import Pt

from resume_tailor.documents.blocks import Span
from resume_tailor.render.docx_common import (
    add_spans,
    set_cell_border,
    set_cell_margins,
    set_indent,
    set_spacing,
)

if TYPE_CHECKING:
    from docx.text.paragraph import Paragraph as DocxParagraph


def test_empty_spans_add_no_runs() -> None:
    paragraph = new_docx().add_paragraph()
    add_spans(paragraph, (Span(""), Span("kept"), Span("")), font="Calibri", size=10)
    assert [run.text for run in paragraph.runs] == ["kept"]


def test_defaults_layer_over_span_emphasis() -> None:
    paragraph = new_docx().add_paragraph()
    add_spans(
        paragraph,
        (Span("a"), Span("b", bold=True)),
        font="Calibri",
        size=10,
        italic=True,
        color="444444",
    )
    assert [(run.bold, run.italic) for run in paragraph.runs] == [(False, True), (True, True)]
    assert str(paragraph.runs[0].font.color.rgb) == "444444"


def test_colour_is_left_alone_when_unset() -> None:
    paragraph = new_docx().add_paragraph()
    add_spans(paragraph, (Span("a"),), font="Calibri", size=10)
    assert paragraph.runs[0].font.color.rgb is None


def _ppr_children(paragraph: DocxParagraph) -> list[str]:
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    properties = paragraph._p.find(f"{namespace}pPr")
    return [child.tag.removeprefix(namespace) for child in properties]


def test_ppr_children_stay_in_schema_order() -> None:
    """w:spacing, w:ind, w:contextualSpacing in that order; out of order, Word calls it corrupt."""
    paragraph = new_docx().add_paragraph("Heading")
    set_spacing(paragraph, before=10, line=13.5)
    set_indent(paragraph, left=0.25, right=0, hanging=0.25)
    children = _ppr_children(paragraph)
    assert children.index("spacing") < children.index("ind") < children.index("contextualSpacing")


def test_spacing_is_absolute_and_lives_above_the_paragraph() -> None:
    paragraph = new_docx().add_paragraph()
    set_spacing(paragraph, before=12, line=13.5)
    fmt = paragraph.paragraph_format
    assert fmt.space_before == Pt(12)
    assert fmt.space_after == Pt(0)
    assert fmt.line_spacing == Pt(13.5)


def test_spacing_disables_contextual_suppression() -> None:
    """List Bullet suppresses the gap between items; the browser has no such rule."""
    paragraph = new_docx().add_paragraph(style="List Bullet")
    set_spacing(paragraph, before=2, line=13.5)
    assert 'w:contextualSpacing w:val="0"' in paragraph._p.xml


def test_spacing_is_idempotent() -> None:
    paragraph = new_docx().add_paragraph(style="List Bullet")
    set_spacing(paragraph, before=2, line=13.5)
    set_spacing(paragraph, before=4, line=13.5)
    assert _ppr_children(paragraph).count("contextualSpacing") == 1
    assert paragraph.paragraph_format.space_before == Pt(4)


def test_cell_margins_are_zeroed() -> None:
    """Word's default table style pads every cell, narrowing the .docx columns vs the HTML."""
    table = new_docx().add_table(rows=1, cols=2)
    set_cell_margins(table, 0)
    xml = table._tbl.xml
    assert "w:tblCellMar" in xml
    assert xml.count('w:w="0"') >= 4


def test_cell_borders_accumulate_rather_than_replace() -> None:
    cell = new_docx().add_table(rows=1, cols=1).rows[0].cells[0]
    set_cell_border(cell, "right")
    set_cell_border(cell, "top", color="FF0000")
    borders = cell._tc.find(qn("w:tcPr")).findall(qn("w:tcBorders"))
    assert len(borders) == 1
    assert {child.tag.split("}")[1] for child in borders[0]} == {"right", "top"}


def test_indent_without_a_hanging_first_line() -> None:
    paragraph = new_docx().add_paragraph()
    set_indent(paragraph, left=0.5, right=0.25)
    assert paragraph.paragraph_format.left_indent == Pt(36)
    assert paragraph.paragraph_format.right_indent == Pt(18)
    assert paragraph.paragraph_format.first_line_indent is None


def test_indent_with_a_hanging_first_line() -> None:
    paragraph = new_docx().add_paragraph()
    set_indent(paragraph, left=0.5, right=0.25, hanging=0.25)
    assert paragraph.paragraph_format.first_line_indent == Pt(-18)
