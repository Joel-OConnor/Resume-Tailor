"""Pin the typography and geometry of both layouts to literal values.

Every other render test compares output against the module constant that produced it, so it moves
in lockstep with any change to that constant and can never catch one. These assertions spell the
numbers out. Changing a constant is then a deliberate two-place edit, and the diff says what the
document will look like afterwards.

Both layouts take their proportions from the user's own "2026 Polished Resume" PDF, set in Arial
at the sizes resume guidance for people and parsers alike recommends (body 10-12pt, headings
14-16pt): the single-column layout carries that typography in the parser-safe shape
``reference/ATS-PLAYBOOK.md`` describes, and the polished layout carries its two-column geometry
as well.
"""

from __future__ import annotations

import re
from typing import TYPE_CHECKING

import pytest
from docx import Document as read_docx
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Inches, Length, Pt

from resume_tailor.render import ats, polished
from resume_tailor.render.html_common import FONT_STACK

if TYPE_CHECKING:
    from collections.abc import Sequence
    from pathlib import Path
    from types import ModuleType

    from docx.table import _Cell
    from docx.text.paragraph import Paragraph as DocxParagraph

    from resume_tailor.documents.blocks import Document

_TWIP = 635  # EMU

# Arial's natural line box as a share of its size, about 1.15: (ascender + descender + line gap)
# over units per em, from its hhea table. Word's "single" line and the browser's "normal" are this.
_ARIAL_LINE = (1854 + 434 + 67) / 2048


# --- the numbers themselves -----------------------------------------------------------------------
@pytest.mark.parametrize(
    ("name", "expected"),
    [
        ("FONT", "Arial"),
        ("BODY_PT", 10.0),
        ("PROSE_PT", 11.0),
        ("CONTACT_PT", 10.0),
        ("SUBTITLE_PT", 12.0),
        ("SECTION_PT", 15.0),
        ("NAME_PT", 28.0),
        ("LINE_PT", 12.0),
        ("PROSE_LINE_PT", 13.5),
        ("SUBTITLE_LINE_PT", 14.25),
        ("SECTION_LINE_PT", 18.0),
        ("NAME_LINE_PT", 33.0),
        ("PAGE_MARGIN_IN", (0.6, 0.7)),
        ("BULLET_HANGING_IN", 0.25),
        ("NOTE_INDENT_IN", 0.05),
        ("BULLET_GLYPH", "\u2022"),
    ],
)
def test_ats_constants(name: str, expected: object) -> None:
    assert getattr(ats, name) == expected


def test_ats_gaps() -> None:
    assert ats.GAP_PT == {
        "name": 0.0,
        "subtitle": 2.25,
        "contact": 2.25,
        "section": 12.0,
        "section-after-name": 37.5,
        "section-after-header": 21.0,
        "entry": 5.25,
        "meta": 1.5,
        "skill": 2.25,
        "bullet": 2.25,
        "bullet-first": 6.75,
        "note": 2.25,
        "para": 5.25,
    }


@pytest.mark.parametrize(
    ("name", "expected"),
    [
        ("FONT", "Arial"),
        ("BODY_PT", 10.0),
        ("PROSE_PT", 11.0),
        ("CONTACT_PT", 10.0),
        ("SECTION_PT", 15.0),
        ("SUBTITLE_PT", 12.0),
        ("NAME_PT", 28.0),
        ("LINE_PT", 12.0),
        ("PROSE_LINE_PT", 13.5),
        ("SUBTITLE_LINE_PT", 14.25),
        ("SECTION_LINE_PT", 18.0),
        ("NAME_LINE_PT", 33.0),
        ("PAGE_MARGIN_IN", 0.2),
        ("SIDEBAR_WIDTH_IN", 2.42),
        ("MAIN_WIDTH_IN", 6.08),
        ("SIDEBAR_INDENT", (0.2, 0.2)),
        ("MAIN_INDENT", (0.2, 0.25)),
        ("BULLET_HANGING_IN", 0.25),
        ("NOTE_INDENT_IN", 0.05),
        ("BULLET_GLYPH", "\u2022"),
    ],
)
def test_polished_constants(name: str, expected: object) -> None:
    assert getattr(polished, name) == expected


def test_polished_gaps() -> None:
    assert polished.GAP_PT == {
        "name": 0.0,
        "contact": 4.5,
        "subtitle": 2.25,
        "section": 9.75,
        "section-rail": 15.75,
        "section-after-name": 37.5,
        "section-after-subtitle": 21.0,
        "entry": 5.25,
        "meta": 1.5,
        "bullet": 2.25,
        "bullet-first": 6.75,
        "note": 2.25,
        "para": 5.25,
        "skill-group": 3.75,
        "skill-item": 1.5,
        "skill-item-first": 3.75,
        "skill-inline": 2.25,
    }


@pytest.mark.parametrize("module", [ats, polished])
def test_the_pdf_fetches_nothing_at_print_time(module: ModuleType) -> None:
    """Arial is a system font on macOS and Windows, so there is no web font to wait for."""
    assert "@import" not in module.CSS
    assert "url(" not in module.CSS


@pytest.mark.parametrize("module", [ats, polished])
def test_the_css_asks_only_for_the_weights_arial_has(module: ModuleType) -> None:
    """Regular and bold, the two a .docx run can say; a 200 or a 600 would have no Word twin."""
    assert set(re.findall(r"font-weight: *(\d+)", module.CSS)) <= {"400", "700"}


_SIZED_LINES = (
    ("BODY_PT", "LINE_PT"),
    ("CONTACT_PT", "LINE_PT"),
    ("PROSE_PT", "PROSE_LINE_PT"),
    ("SUBTITLE_PT", "SUBTITLE_LINE_PT"),
    ("SECTION_PT", "SECTION_LINE_PT"),
    ("NAME_PT", "NAME_LINE_PT"),
)


@pytest.mark.parametrize("module", [ats, polished])
@pytest.mark.parametrize(("size", "line"), _SIZED_LINES)
def test_every_line_box_holds_a_line_of_arial(module: ModuleType, size: str, line: str) -> None:
    """Word's "at least" widens a box smaller than the font's own; the browser does not."""
    assert getattr(module, line) >= getattr(module, size) * _ARIAL_LINE


_LINE_CONSTANTS = (
    "LINE_PT",
    "PROSE_LINE_PT",
    "SUBTITLE_LINE_PT",
    "SECTION_LINE_PT",
    "NAME_LINE_PT",
)


def test_every_vertical_measure_is_a_whole_css_pixel() -> None:
    """The browser snaps to whole pixels (0.75pt); a value that already is one cannot drift."""
    measures = [
        *(getattr(module, name) for module in (ats, polished) for name in _LINE_CONSTANTS),
        *ats.GAP_PT.values(),
        *polished.GAP_PT.values(),
    ]
    assert all((value / 0.75).is_integer() for value in measures)


def test_the_columns_fill_the_page() -> None:
    """Full bleed is the whole point of the design; the rule sits where the split falls."""
    assert polished.SIDEBAR_WIDTH_IN + polished.MAIN_WIDTH_IN == 8.5


def test_the_font_stack_is_arial_then_its_metric_twin() -> None:
    assert FONT_STACK == "Arial, Helvetica, sans-serif"


# --- the numbers as they reach the .docx ----------------------------------------------------------
# One paragraph of each block kind in the fixture, found by its opening words, and the run to read
# in it: (opening, run, size, bold, italic, line box, gap above), sizes and spacing in points.
_Typography = tuple[str, int, float, bool, bool, float, float]

_ATS_BLOCKS: tuple[_Typography, ...] = (
    # A large name reads as prominent without the weight, and headings are regular too.
    ("Ada Lovelace", 0, 28, False, False, 33, 0),
    ("Principal Engineer", 0, 12, False, False, 14.25, 2.25),
    ("ada@example.com", 0, 10, False, False, 12, 2.25),
    ("Summary", 0, 15, False, False, 18, 21),  # follows the contact line
    ("Engineer who writes", 0, 11, False, False, 13.5, 5.25),
    ("Skills", 0, 15, False, False, 18, 12),  # follows prose
    ("Languages", 0, 10, True, False, 12, 2.25),  # a skill line's label...
    ("Languages", 1, 10, False, False, 12, 2.25),  # ...and its items
    ("Analytical Engine", 0, 10, True, False, 12, 5.25),  # an entry with no bold span
    ("London, UK", 0, 10, False, True, 12, 1.5),  # its dates line
    ("Algorithm Design", 1, 10, False, False, 12, 6.75),  # the first bullet of a run...
    ("Corresponded", 0, 10, False, False, 12, 2.25),  # ...and the next
    ("Tech Stack", 0, 10, False, True, 12, 2.25),  # the closing note
)

_POLISHED_MAIN: tuple[_Typography, ...] = (
    ("Ada Lovelace", 0, 28, False, False, 33, 0),
    ("Principal Engineer", 0, 12, False, False, 14.25, 2.25),
    ("Summary", 0, 15, False, False, 18, 21),  # follows the target title
    ("Engineer who writes", 0, 11, False, False, 13.5, 5.25),
    ("Experience", 0, 15, False, False, 18, 9.75),  # follows prose
    ("Analytical Engine", 0, 10, True, False, 12, 5.25),
    ("London, UK", 0, 10, False, True, 12, 1.5),
    ("Algorithm Design", 1, 10, False, False, 12, 6.75),
    ("Corresponded", 0, 10, False, False, 12, 2.25),
    ("Tech Stack", 0, 10, False, True, 12, 2.25),
)

_POLISHED_RAIL: tuple[_Typography, ...] = (
    ("ada@example.com", 0, 10, False, False, 12, 0),  # opens the column
    ("(555) 010-0100", 0, 10, False, False, 12, 4.5),
    ("Skills", 0, 15, False, False, 18, 15.75),
    ("Languages", 0, 10, True, False, 12, 3.75),  # a skill group's label...
    ("Analytical Notation", 0, 10, False, False, 12, 3.75),  # ...its first item...
    ("Mathematics", 0, 10, False, False, 12, 1.5),  # ...and the next
    ("Private tuition", 0, 10, True, False, 12, 5.25),
    ("De Morgan", 0, 10, False, True, 12, 1.5),
    ("Fellow of the", 0, 10, False, False, 12, 3.75),  # a rail bullet under a heading
)


def _docx_points(length: object) -> float | None:
    """Return a python-docx length in points, or ``None`` for anything else (unset, a multiple)."""
    return float(length.pt) if isinstance(length, Length) else None


def _typography(
    paragraphs: Sequence[DocxParagraph], expected: Sequence[_Typography]
) -> list[tuple[object, ...]]:
    """Read back each expected row from the first paragraph that opens with its words."""
    found: list[tuple[object, ...]] = []
    for opening, index, *_ in expected:
        paragraph = next(p for p in paragraphs if p.text.startswith(opening))
        run = paragraph.runs[index]
        fmt = paragraph.paragraph_format
        found.append(
            (
                opening,
                index,
                _docx_points(run.font.size),
                run.bold,
                run.italic,
                _docx_points(fmt.line_spacing),
                _docx_points(fmt.space_before),
            )
        )
    return found


def _insets(cell: _Cell) -> set[tuple[float | None, ...]]:
    """Every (left, right, first-line) indent the paragraphs of one column use, in inches."""
    return {
        tuple(
            None if length is None else length.inches
            for length in (fmt.left_indent, fmt.right_indent, fmt.first_line_indent)
        )
        for fmt in (p.paragraph_format for p in cell.paragraphs)
    }


def test_ats_page_geometry(resume: Document, tmp_path: Path) -> None:
    out = tmp_path / "resume.docx"
    ats.render_docx(resume, out)
    section = read_docx(str(out)).sections[0]
    assert section.top_margin == Pt(0.6 * 72)
    assert section.bottom_margin == Pt(0.6 * 72)
    assert section.left_margin == Pt(0.7 * 72)
    assert section.right_margin == Pt(0.7 * 72)


def test_ats_typography_reaches_the_docx(resume: Document, tmp_path: Path) -> None:
    out = tmp_path / "resume.docx"
    ats.render_docx(resume, out)
    paragraphs = read_docx(str(out)).paragraphs
    assert _typography(paragraphs, _ATS_BLOCKS) == list(_ATS_BLOCKS)
    # Every gap sits above a block and none below: Word adds the two where CSS collapses them.
    assert {p.paragraph_format.space_after for p in paragraphs} == {Pt(0)}


def test_polished_page_geometry(resume: Document, tmp_path: Path) -> None:
    out = tmp_path / "resume-polished.docx"
    polished.render_docx(resume, out)
    document = read_docx(str(out))
    section = document.sections[0]
    assert section.top_margin == Inches(0.2)
    assert section.bottom_margin == Inches(0.2)
    assert section.left_margin == Inches(0)
    assert section.right_margin == Inches(0)
    columns = document.tables[0].columns
    assert abs(columns[0].width - Inches(2.42)) <= _TWIP
    assert abs(columns[1].width - Inches(6.08)) <= _TWIP


def test_polished_typography_reaches_the_docx(resume: Document, tmp_path: Path) -> None:
    out = tmp_path / "resume-polished.docx"
    polished.render_docx(resume, out)
    rail, main = read_docx(str(out)).tables[0].rows[0].cells
    assert _typography(main.paragraphs, _POLISHED_MAIN) == list(_POLISHED_MAIN)
    assert _typography(rail.paragraphs, _POLISHED_RAIL) == list(_POLISHED_RAIL)
    after = {p.paragraph_format.space_after for cell in (rail, main) for p in cell.paragraphs}
    assert after == {Pt(0)}


def test_a_polished_skill_line_in_the_main_column_bolds_only_its_label(
    resume: Document, tmp_path: Path
) -> None:
    out = tmp_path / "resume-polished.docx"
    polished.render_docx(resume, out, sidebar_sections=("summary",))
    main = read_docx(str(out)).tables[0].rows[0].cells[1]
    inline: tuple[_Typography, ...] = (
        ("Languages", 0, 10, True, False, 12, 2.25),
        ("Languages", 1, 10, False, False, 12, 2.25),
    )
    assert _typography(main.paragraphs, inline) == list(inline)


def test_each_polished_column_insets_its_text_by_its_own_padding(
    resume: Document, tmp_path: Path
) -> None:
    """The .docx twin of the CSS padding: 0.2in either side of the rail, 0.2in and 0.25in main."""
    out = tmp_path / "resume-polished.docx"
    polished.render_docx(resume, out)
    rail, main = read_docx(str(out)).tables[0].rows[0].cells
    # Plain text, then a bullet hanging 0.25in past its glyph, then a note inset 0.05in more.
    assert _insets(rail) == {(0.2, 0.2, None), (0.45, 0.2, -0.25)}
    assert _insets(main) == {(0.2, 0.25, None), (0.45, 0.25, -0.25), (0.25, 0.25, None)}


def test_every_polished_run_is_arial(resume: Document, tmp_path: Path) -> None:
    """One family in every run of both cells, so Word never substitutes a face it lacks."""
    out = tmp_path / "resume-polished.docx"
    polished.render_docx(resume, out)
    cells = read_docx(str(out)).tables[0].rows[0].cells
    assert {run.font.name for cell in cells for p in cell.paragraphs for run in p.runs} == {"Arial"}


def test_polished_zeroes_the_table_cell_margins(resume: Document, tmp_path: Path) -> None:
    """Word's default 0.075in cell padding would narrow the .docx columns against the PDF."""
    out = tmp_path / "resume-polished.docx"
    polished.render_docx(resume, out)
    xml = read_docx(str(out)).tables[0]._tbl.xml
    assert "w:tblCellMar" in xml
    assert xml.count('w:w="0" w:type="dxa"') >= 4


def test_the_divider_is_one_point(resume: Document, tmp_path: Path) -> None:
    """w:sz is in eighths of a point, so 8 is 1pt — and the CSS must say the same."""
    out = tmp_path / "resume-polished.docx"
    polished.render_docx(resume, out)
    rail = read_docx(str(out)).tables[0].rows[0].cells[0]
    assert 'w:sz="8"' in rail._tc.xml
    assert "border-right: 1pt solid #000" in polished.render_html(resume)


def test_line_spacing_is_at_least_never_exact(resume: Document, tmp_path: Path) -> None:
    """EXACTLY clips any glyph taller than the line box — Word would silently truncate text."""
    out = tmp_path / "resume.docx"
    ats.render_docx(resume, out)
    rules = {p.paragraph_format.line_spacing_rule for p in read_docx(str(out)).paragraphs}
    assert rules == {WD_LINE_SPACING.AT_LEAST}


# --- the numbers as they reach the HTML (and so the PDF) ------------------------------------------
# The same literals the .docx is held to above, so the two renders agree by construction: a change
# to one side alone fails here or there.
_CSS_COMMENT = re.compile(r"/\*.*?\*/", re.DOTALL)
_CSS_RULE = re.compile(r"([^{}]+)\{([^}]*)\}")


def _css_rules(css: str) -> dict[str, dict[str, str]]:
    """Read a flat stylesheet into ``{selector: {property: value}}``, one key per selector."""
    rules: dict[str, dict[str, str]] = {}
    for selectors, body in _CSS_RULE.findall(_CSS_COMMENT.sub("", css)):
        declarations = {
            name.strip(): value.strip()
            for name, _, value in (item.partition(":") for item in body.split(";"))
            if name.strip()
        }
        for selector in selectors.split(","):
            rules.setdefault(" ".join(selector.split()), {}).update(declarations)
    return rules


def _css_points(value: str | None) -> float | None:
    """Read a CSS length in points, ``12.0pt`` as 12.0; ``None`` where no rule sets it."""
    return None if value is None else float(value.removesuffix("pt"))


# (selector, property) -> points.
_ATS_CSS_POINTS: dict[tuple[str, str], float] = {
    ("body", "font-size"): 10,
    ("body", "line-height"): 12,
    ("h1", "font-size"): 28,
    ("h1", "line-height"): 33,
    ("h1", "margin-top"): 0,
    (".subtitle", "font-size"): 12,
    (".subtitle", "line-height"): 14.25,
    (".subtitle", "margin-top"): 2.25,
    (".contact", "font-size"): 10,
    (".contact", "margin-top"): 2.25,
    ("h2", "font-size"): 15,
    ("h2", "line-height"): 18,
    ("h2", "margin-top"): 12,
    ("h1 + h2", "margin-top"): 37.5,
    (".subtitle + h2", "margin-top"): 21,
    (".contact + h2", "margin-top"): 21,
    ("h3", "font-size"): 10,
    ("h3", "margin-top"): 5.25,
    (".meta", "margin-top"): 1.5,
    (".skill", "margin-top"): 2.25,
    ("p", "font-size"): 11,
    ("p", "line-height"): 13.5,
    ("p", "margin-top"): 5.25,
    ("p.note", "font-size"): 10,
    ("p.note", "line-height"): 12,
    ("p.note", "margin-top"): 2.25,
    ("li", "margin-top"): 2.25,
    ("li:first-child", "margin-top"): 6.75,
    ("body > :first-child", "margin-top"): 0,
}

# (selector, property) -> the declared value, verbatim.
_ATS_CSS_VALUES: dict[tuple[str, str], str] = {
    ("@page", "size"): "Letter",
    ("@page", "margin"): "0.6in 0.7in",
    # A browser bolds h1-h3 by default; the .docx name and headings are regular.
    ("h1", "font-weight"): "400",
    ("h2", "font-weight"): "400",
    ("h3", "font-weight"): "400",
    (".meta", "font-style"): "italic",
    ("p.note", "padding-left"): "0.05in",
    ("ul", "list-style"): "none",
    ("li", "padding-left"): "0.25in",
}

_POLISHED_CSS_POINTS: dict[tuple[str, str], float] = {
    ("body", "font-size"): 10,
    ("body", "line-height"): 12,
    ("h1", "font-size"): 28,
    ("h1", "line-height"): 33,
    ("h1", "margin-top"): 0,
    ("h2", "font-size"): 15,
    ("h2", "line-height"): 18,
    ("h2", "margin-top"): 9.75,
    (".rail h2", "margin-top"): 15.75,
    ("h1 + h2", "margin-top"): 37.5,
    (".subtitle + h2", "margin-top"): 21,
    ("h3", "font-size"): 10,
    ("h3", "margin-top"): 5.25,
    (".subtitle", "font-size"): 12,
    (".subtitle", "line-height"): 14.25,
    (".subtitle", "margin-top"): 2.25,
    (".contact", "font-size"): 10,
    (".contact", "margin-top"): 4.5,
    (".meta", "margin-top"): 1.5,
    (".skill-group", "margin-top"): 3.75,
    (".skill-inline", "margin-top"): 2.25,
    ("p", "font-size"): 11,
    ("p", "line-height"): 13.5,
    ("p", "margin-top"): 5.25,
    ("p.note", "font-size"): 10,
    ("p.note", "line-height"): 12,
    ("p.note", "margin-top"): 2.25,
    ("li", "margin-top"): 2.25,
    ("li:first-child", "margin-top"): 6.75,
    (".rail li", "margin-top"): 1.5,
    (".rail li:first-child", "margin-top"): 3.75,
    (".rail > :first-child", "margin-top"): 0,
    (".main > :first-child", "margin-top"): 0,
}

_POLISHED_CSS_VALUES: dict[tuple[str, str], str] = {
    ("@page", "size"): "Letter",
    ("@page", "margin"): "0.2in 0",
    (".rail", "width"): "2.42in",
    (".rail", "flex"): "0 0 2.42in",
    (".rail", "padding"): "0 0.2in 0 0.2in",
    (".main", "width"): "6.08in",
    (".main", "flex"): "1 1 6.08in",
    (".main", "padding"): "0 0.25in 0 0.2in",
    ("h1", "font-weight"): "400",
    ("h2", "font-weight"): "400",
    ("h3", "font-weight"): "400",
    (".skill-group", "font-weight"): "700",  # the rail label, bold in the .docx too
    (".meta", "font-style"): "italic",
    ("p.note", "padding-left"): "0.05in",
    ("ul", "list-style"): "none",
    ("li", "padding-left"): "0.25in",
}


@pytest.mark.parametrize(
    ("module", "points", "values"),
    [
        (ats, _ATS_CSS_POINTS, _ATS_CSS_VALUES),
        (polished, _POLISHED_CSS_POINTS, _POLISHED_CSS_VALUES),
    ],
    ids=["ats", "polished"],
)
def test_the_css_sets_the_numbers_the_docx_does(
    module: ModuleType,
    points: dict[tuple[str, str], float],
    values: dict[tuple[str, str], str],
) -> None:
    rules = _css_rules(module.CSS)
    declared = {key: rules.get(key[0], {}).get(key[1]) for key in (*points, *values)}
    assert {key: _css_points(declared[key]) for key in points} == points
    assert {key: declared[key] for key in values} == values
