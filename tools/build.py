#!/usr/bin/env python3
"""Render a resume/cover-letter Markdown file into an ATS-safe .docx and a .pdf.

Usage:
    python tools/build.py path/to/resume.md [--out-dir DIR]

The Markdown must follow templates/resume.md (or templates/cover-letter.md): single column, standard
headings, no tables/images. The .docx is built with python-docx (the format resume screeners parse
most reliably); the .pdf is produced by printing a clean HTML render via headless Google Chrome, so
its text stays selectable (also ATS-readable). Both come out next to the source unless --out-dir.
"""

from __future__ import annotations

import argparse
import html
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

# --- parse the constrained resume/letter Markdown into a flat list of blocks ----------------------
# Each block is a (kind, payload) tuple:
#   ('name', str) ('header', str) ('h2', str) ('h3', str) ('meta', str)
#   ('skill', (label, items)) ('bullet', str) ('para', str)

_COMMENT = re.compile(r"<!--.*?-->", re.DOTALL)
_SKILL = re.compile(r"^\*\*(.+?):\*\*\s*(.*)$")


def parse(md: str) -> list[tuple]:
    md = _COMMENT.sub("", md)
    lines = md.splitlines()
    blocks: list[tuple] = []

    # locate the name line ("# Name") and skip anything before it
    i = 0
    while i < len(lines) and not re.match(r"^#\s+\S", lines[i]):
        i += 1
    if i >= len(lines):
        raise SystemExit("no '# Name' line found — start the file with '# Full Name'")
    blocks.append(("name", lines[i].lstrip("# ").strip()))
    i += 1

    # header block: consecutive non-blank lines until a blank line or a section heading
    while i < len(lines) and lines[i].strip() and not lines[i].startswith("## "):
        blocks.append(("header", lines[i].strip()))
        i += 1

    # body
    para: list[str] = []
    after_h3 = False

    def flush():
        nonlocal para
        if para:
            blocks.append(("para", " ".join(para)))
            para = []

    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            flush()
            after_h3 = False
        elif stripped.startswith("## "):
            flush()
            blocks.append(("h2", stripped[3:].strip()))
            after_h3 = False
        elif stripped.startswith("### "):
            flush()
            blocks.append(("h3", stripped[4:].strip()))
            after_h3 = True
        elif stripped.startswith(("- ", "* ")):
            flush()
            blocks.append(("bullet", stripped[2:].strip()))
            after_h3 = False
        elif _SKILL.match(stripped):
            flush()
            m = _SKILL.match(stripped)
            blocks.append(("skill", (m.group(1).strip(), m.group(2).strip())))
            after_h3 = False
        elif stripped == "---":
            flush()
            after_h3 = False
        else:
            # a plain line: the first one right after ### is the dates/location meta line
            if after_h3 and not para:
                blocks.append(("meta", stripped))
                after_h3 = False
            else:
                para.append(stripped)
        i += 1
    flush()
    return blocks


# --- .docx rendering ------------------------------------------------------------------------------
_BOLD = re.compile(r"(\*\*.+?\*\*)")


def _runs(paragraph, text, *, bold=False, italic=False, size=None, color=None):
    """Add text to a paragraph, honoring inline **bold** segments."""
    for part in _BOLD.split(text):
        if not part:
            continue
        b, t = bold, part
        if part.startswith("**") and part.endswith("**"):
            b, t = True, part[2:-2]
        run = paragraph.add_run(t)
        run.bold = b
        run.italic = italic
        if size:
            run.font.size = size
        if color:
            run.font.color.rgb = color


def _bottom_border(paragraph):
    """A thin rule under a section heading (a paragraph border — safe for ATS, not a table)."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "2"), ("w:color", "999999")):
        bottom.set(qn(k), v)
    pBdr.append(bottom)
    pPr.append(pBdr)


def render_docx(blocks: list[tuple], out: Path) -> None:
    doc = Document()
    normal = doc.styles["Normal"].font
    normal.name = "Calibri"
    normal.size = Pt(10.5)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Pt(43)  # ~0.6in
        section.left_margin = section.right_margin = Pt(50)  # ~0.7in

    for kind, payload in blocks:
        if kind == "name":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            _runs(p, payload, bold=True, size=Pt(20))
        elif kind == "header":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            _runs(p, payload, size=Pt(10), color=RGBColor(0x44, 0x44, 0x44))
        elif kind == "h2":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            _runs(p, payload.upper(), bold=True, size=Pt(12))
            _bottom_border(p)
        elif kind == "h3":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(7)
            p.paragraph_format.space_after = Pt(0)
            _runs(p, payload, bold=True, size=Pt(11))
        elif kind == "meta":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            _runs(p, payload, italic=True, size=Pt(10), color=RGBColor(0x44, 0x44, 0x44))
        elif kind == "skill":
            label, items = payload
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            _runs(p, f"{label}: ", bold=True, size=Pt(10))
            _runs(p, items, size=Pt(10))
        elif kind == "bullet":
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(1)
            _runs(p, payload, size=Pt(10))
        elif kind == "para":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            _runs(p, payload, size=Pt(10.5))

    doc.save(str(out))


# --- .pdf rendering (clean HTML → headless Chrome print-to-pdf) ------------------------------------
_CSS = """
@page { size: Letter; margin: 0.6in 0.7in; }
* { box-sizing: border-box; }
body { font-family: Calibri, Helvetica, Arial, sans-serif; font-size: 10.5pt; line-height: 1.3;
       color: #111; margin: 0; }
h1 { font-size: 20pt; margin: 0 0 2pt 0; }
.header { color: #444; font-size: 10pt; margin: 0 0 1pt 0; }
h2 { font-size: 12pt; text-transform: uppercase; border-bottom: 1px solid #999;
     margin: 12pt 0 4pt 0; padding-bottom: 2pt; }
h3 { font-size: 11pt; margin: 8pt 0 0 0; }
.meta { color: #444; font-style: italic; font-size: 10pt; margin: 0 0 2pt 0; }
.skill { margin: 0 0 1pt 0; }
p { margin: 0 0 5pt 0; }
ul { margin: 2pt 0 4pt 0; padding-left: 18pt; }
li { margin: 0 0 1pt 0; }
"""


def _inline_html(text: str) -> str:
    return _BOLD.sub(lambda m: f"<strong>{html.escape(m.group(1)[2:-2])}</strong>", html.escape(text))


def render_html(blocks: list[tuple]) -> str:
    out: list[str] = []
    open_ul = False

    def close_ul():
        nonlocal open_ul
        if open_ul:
            out.append("</ul>")
            open_ul = False

    for kind, payload in blocks:
        if kind != "bullet":
            close_ul()
        if kind == "name":
            out.append(f"<h1>{_inline_html(payload)}</h1>")
        elif kind == "header":
            out.append(f'<div class="header">{_inline_html(payload)}</div>')
        elif kind == "h2":
            out.append(f"<h2>{_inline_html(payload)}</h2>")
        elif kind == "h3":
            out.append(f"<h3>{_inline_html(payload)}</h3>")
        elif kind == "meta":
            out.append(f'<div class="meta">{_inline_html(payload)}</div>')
        elif kind == "skill":
            label, items = payload
            out.append(f'<div class="skill"><strong>{html.escape(label)}:</strong> {_inline_html(items)}</div>')
        elif kind == "bullet":
            if not open_ul:
                out.append("<ul>")
                open_ul = True
            out.append(f"<li>{_inline_html(payload)}</li>")
        elif kind == "para":
            out.append(f"<p>{_inline_html(payload)}</p>")
    close_ul()
    return f"<!doctype html><html><head><meta charset='utf-8'><style>{_CSS}</style></head><body>{''.join(out)}</body></html>"


def _find_chrome() -> str | None:
    candidates = [
        "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
        "/Applications/Chromium.app/Contents/MacOS/Chromium",
        "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
        "/Applications/Brave Browser.app/Contents/MacOS/Brave Browser",
    ]
    for c in candidates:
        if os.path.exists(c):
            return c
    return shutil.which("google-chrome") or shutil.which("chromium") or shutil.which("chrome")


def render_pdf(blocks: list[tuple], out: Path) -> bool:
    chrome = _find_chrome()
    html_doc = render_html(blocks)
    with tempfile.NamedTemporaryFile("w", suffix=".html", delete=False, encoding="utf-8") as f:
        f.write(html_doc)
        html_path = f.name
    # keep the HTML next to the output too, as a fallback the user can print by hand
    out.with_suffix(".html").write_text(html_doc, encoding="utf-8")
    if not chrome:
        print(f"  · no Chrome found — wrote {out.with_suffix('.html').name}; open it and Cmd+P → Save as PDF")
        return False
    for headless in ("--headless=new", "--headless"):
        try:
            r = subprocess.run(
                [chrome, headless, "--disable-gpu", "--no-pdf-header-footer",
                 f"--print-to-pdf={out}", f"file://{html_path}"],
                capture_output=True, timeout=60,
            )
        except Exception as exc:  # noqa: BLE001
            print(f"  · Chrome PDF failed ({exc})")
            return False
        if out.exists() and out.stat().st_size > 0:
            os.unlink(html_path)
            out.with_suffix(".html").unlink(missing_ok=True)
            return True
    print(f"  · Chrome could not produce a PDF; wrote {out.with_suffix('.html').name} to print by hand")
    return False


# --- CLI ------------------------------------------------------------------------------------------
def build(md_path: Path, out_dir: Path | None = None) -> None:
    blocks = parse(md_path.read_text(encoding="utf-8"))
    out_dir = out_dir or md_path.parent
    out_dir.mkdir(parents=True, exist_ok=True)
    docx_path = out_dir / (md_path.stem + ".docx")
    pdf_path = out_dir / (md_path.stem + ".pdf")
    render_docx(blocks, docx_path)
    print(f"  ✓ {docx_path}")
    if render_pdf(blocks, pdf_path):
        print(f"  ✓ {pdf_path}")


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Render resume/cover-letter Markdown to .docx + .pdf")
    ap.add_argument("files", nargs="+", help="markdown file(s) to render")
    ap.add_argument("--out-dir", default=None, help="write outputs here (default: next to the source)")
    args = ap.parse_args(argv)
    out_dir = Path(args.out_dir) if args.out_dir else None
    for f in args.files:
        p = Path(f)
        if not p.exists():
            print(f"  ! not found: {p}", file=sys.stderr)
            continue
        print(f"{p}:")
        build(p, out_dir)
    return 0


if __name__ == "__main__":
    sys.exit(main())
