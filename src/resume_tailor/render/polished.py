"""The polished layout: a two-column, typographic resume for human readers.

This reproduces the arrangement of the user's own "2026 Polished Resume": a narrow left rail
holding contact details, skills, and education, separated by a hairline rule from a wide right
column holding the name, summary, and experience. The geometry was measured from that PDF: 0.2in
page and column margins, a 12pt line pitch for body text and 13.5pt for the summary, and a bullet
glyph with a 0.25in hanging indent on every accomplishment and on every skill in the rail. The
face is Arial throughout, in the only two weights it has: regular for the name, contact lines,
and headings, bold for the job title in a role heading (see :mod:`resume_tailor.render.emphasis`),
lead-ins, and rail labels.

It is **not** ATS-safe: the two columns are a table, and table layouts get scrambled or dropped
by resume parsers. That is why it is opt-in (``--layout polished``): the default export is
:mod:`resume_tailor.render.ats`, the same typography in one column, which any parser can read.

Arial ships with macOS and Windows, so the PDF printed from the HTML fetches nothing at print time
and Word draws the ``.docx`` in the same face the PDF shows.
"""

from __future__ import annotations

from typing import TYPE_CHECKING, assert_never, cast

from docx import Document as new_docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt

from resume_tailor.documents.blocks import (
    Bullet,
    Entry,
    HeaderLine,
    Meta,
    Name,
    Paragraph,
    Section,
    SectionGroup,
    SkillLine,
    Span,
    is_contact_line,
    is_note,
    title_of,
)
from resume_tailor.render.docx_common import (
    add_spans,
    set_cell_border,
    set_cell_margins,
    set_indent,
    set_properties,
    set_spacing,
)
from resume_tailor.render.emphasis import entry_spans
from resume_tailor.render.html_common import FONT_STACK, page, spans_to_html

if TYPE_CHECKING:
    from pathlib import Path

    from docx.table import _Cell
    from docx.text.paragraph import Paragraph as DocxParagraph

    from resume_tailor.documents.blocks import Block, Document, GroupBlock

__all__ = [
    "DEFAULT_SIDEBAR_SECTIONS",
    "render_docx",
    "render_html",
    "split_columns",
]

#: Sections that belong in the left rail. Matched case-insensitively against ``## `` headings.
DEFAULT_SIDEBAR_SECTIONS: tuple[str, ...] = (
    "skills",
    "technical core",
    "core competencies",
    "education",
    "certifications",
)

FONT = "Arial"
BODY_PT = 10.0
PROSE_PT = 11.0
CONTACT_PT = 10.0
SECTION_PT = 15.0
SUBTITLE_PT = 12.0
NAME_PT = 28.0

# Line boxes, absolute so Word and the browser agree (see docx_common.set_spacing). Each is at
# least Arial's natural box at that size, 1.15 times it, so Word's "at least" never has to widen
# one and the two renderers stack lines identically. Every vertical measure here is a multiple of
# 0.75pt, one CSS pixel: the browser snaps line boxes and margins to whole pixels, Word does not,
# and a value that is already whole keeps the PDF and the .docx from drifting a fraction per line.
LINE_PT = 12.0
PROSE_LINE_PT = 13.5
SUBTITLE_LINE_PT = 14.25
SECTION_LINE_PT = 18.0
NAME_LINE_PT = 33.0

# The gap ABOVE each block. It lives in exactly one property on each side, and a block that opens
# its column gets none at all: both columns start at the page margin.
GAP_PT = {
    "name": 0.0,
    "contact": 4.5,
    "subtitle": 2.25,
    "section": 9.75,
    "section-rail": 15.75,
    "section-after-name": 37.5,
    "section-after-subtitle": 21.0,
    "entry": 5.25,
    "meta": 1.5,
    "bullet": 2.25,
    "bullet-first": 6.75,
    "note": 2.25,
    "para": 5.25,
    "skill-group": 3.75,
    "skill-item": 1.5,
    "skill-item-first": 3.75,
    "skill-inline": 2.25,
}

PAGE_MARGIN_IN = 0.2
SIDEBAR_WIDTH_IN = 2.42
MAIN_WIDTH_IN = 6.08
SIDEBAR_INDENT = (0.2, 0.2)
MAIN_INDENT = (0.2, 0.25)
BULLET_HANGING_IN = 0.25
NOTE_INDENT_IN = 0.05
BULLET_GLYPH = "•"


def _split_contact(spans: tuple[Span, ...]) -> list[tuple[Span, ...]]:
    """Split a ``a | b | c`` contact line into one stacked entry per item."""
    entries: list[tuple[Span, ...]] = []
    current: list[Span] = []
    for span in spans:
        pieces = span.text.split("|")
        for index, piece in enumerate(pieces):
            if index:
                if current:
                    entries.append(tuple(current))
                current = []
            if text := piece.strip():
                current.append(Span(text, bold=span.bold, italic=span.italic))
    if current:
        entries.append(tuple(current))
    return entries


def split_columns(
    document: Document, sidebar_sections: tuple[str, ...] = DEFAULT_SIDEBAR_SECTIONS
) -> tuple[list[SectionGroup], list[SectionGroup]]:
    """Split ``document`` into (sidebar groups, main groups).

    The preamble is divided by kind rather than by section: the contact line goes to the rail,
    while the name and any target-title line head the main column.
    """
    wanted = {name.casefold() for name in sidebar_sections}
    sidebar: list[SectionGroup] = []
    main: list[SectionGroup] = []
    for group in document.groups():
        if group.title:
            (sidebar if group.key in wanted else main).append(group)
        else:
            preamble_sidebar: list[GroupBlock] = []
            preamble_main: list[GroupBlock] = []
            for block in group.blocks:
                if isinstance(block, HeaderLine) and is_contact_line(block.spans):
                    preamble_sidebar += [HeaderLine(entry) for entry in _split_contact(block.spans)]
                else:
                    preamble_main.append(block)
            sidebar.append(SectionGroup("", tuple(preamble_sidebar)))
            main.append(SectionGroup("", tuple(preamble_main)))
    return sidebar, main


# --- the gaps, by what comes before ---------------------------------------------------------------
def _gap(kind: str, previous: Block | None) -> float:
    """Return the space above a block: its kind's gap, or none when it opens the column."""
    return 0.0 if previous is None else GAP_PT[kind]


def _section_kind(previous: Block | None, *, sidebar: bool) -> str:
    """Return the gap a heading takes: more in the rail, most of all directly under the name."""
    if sidebar:
        return "section-rail"
    if isinstance(previous, Name):
        return "section-after-name"
    if isinstance(previous, HeaderLine):
        return "section-after-subtitle"
    return "section"


def _bullet_kind(previous: Block | None, *, sidebar: bool) -> str:
    """Return the gap a bullet takes: the first of a run stands off, the rest sit tighter."""
    base = "skill-item" if sidebar else "bullet"
    return base if isinstance(previous, Bullet) else f"{base}-first"


# --- .docx ----------------------------------------------------------------------------------------
def render_docx(
    document: Document,
    out: Path,
    sidebar_sections: tuple[str, ...] = DEFAULT_SIDEBAR_SECTIONS,
) -> None:
    """Write ``document`` as the two-column polished ``.docx``."""
    docx = new_docx()
    # Word measures a line by its paragraph mark too, and the mark takes Normal's font. Left at
    # python-docx's default it would be Cambria 11pt, a taller line than the PDF draws.
    normal = docx.styles["Normal"].font
    normal.name = FONT
    normal.size = Pt(BODY_PT)
    for section in docx.sections:
        section.top_margin = section.bottom_margin = Inches(PAGE_MARGIN_IN)
        section.left_margin = section.right_margin = Inches(0)

    table = docx.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    rail, main = table.rows[0].cells
    for cell, width in ((rail, SIDEBAR_WIDTH_IN), (main, MAIN_WIDTH_IN)):
        cell.width = Inches(width)
    for column, width in zip(table.columns, (SIDEBAR_WIDTH_IN, MAIN_WIDTH_IN), strict=True):
        column.width = Inches(width)
    set_cell_margins(table, 0)
    set_cell_border(rail, "right")

    sidebar_groups, main_groups = split_columns(document, sidebar_sections)
    _fill_cell(rail, sidebar_groups, indent=SIDEBAR_INDENT, sidebar=True)
    _fill_cell(main, main_groups, indent=MAIN_INDENT, sidebar=False)
    set_properties(docx, document)
    docx.save(str(out))


def _fill_cell(
    cell: _Cell,
    groups: list[SectionGroup],
    *,
    indent: tuple[float, float],
    sidebar: bool,
) -> None:
    """Render ``groups`` into a table cell, dropping the empty placeholder paragraph.

    The placeholder only goes once there is something to replace it with: a ``<w:tc>`` with no
    block-level child is invalid per ECMA-376, and Word refuses to open the file. That happens
    for real input — a document with nothing routed to the rail, such as a cover letter with no
    contact line.
    """
    placeholder = cell.paragraphs[0]
    previous: Block | None = None
    for group in groups:
        blocks: list[Block] = [Section(group.title)] if group.title else []
        blocks += group.blocks
        for block in blocks:
            _add(cell, block, indent=indent, sidebar=sidebar, previous=previous)
            previous = block
    if len(cell.paragraphs) > 1:
        cell._tc.remove(placeholder._p)  # noqa: SLF001


def _paragraph(  # noqa: PLR0913 - each argument is one independent paragraph property
    cell: _Cell,
    indent: tuple[float, float],
    before: float,
    *,
    hanging: float = 0.0,
    style: str | None = None,
    line: float = LINE_PT,
) -> DocxParagraph:
    paragraph = cast("DocxParagraph", cell.add_paragraph(style=style))
    set_indent(paragraph, left=indent[0] + hanging, right=indent[1], hanging=hanging)
    set_spacing(paragraph, before=before, line=line)
    return paragraph


def _add(  # noqa: C901 - flat dispatch over the block union
    cell: _Cell,
    block: Block,
    *,
    indent: tuple[float, float],
    sidebar: bool,
    previous: Block | None,
) -> None:
    match block:
        case Name(spans):
            paragraph = _paragraph(cell, indent, _gap("name", previous), line=NAME_LINE_PT)
            add_spans(paragraph, spans, font=FONT, size=NAME_PT)
        case HeaderLine(spans) if sidebar:
            paragraph = _paragraph(cell, indent, _gap("contact", previous))
            add_spans(paragraph, spans, font=FONT, size=CONTACT_PT)
        case HeaderLine(spans):
            before = _gap("subtitle", previous)
            paragraph = _paragraph(cell, indent, before, line=SUBTITLE_LINE_PT)
            add_spans(paragraph, spans, font=FONT, size=SUBTITLE_PT)
        case Section(title):
            before = _gap(_section_kind(previous, sidebar=sidebar), previous)
            paragraph = _paragraph(cell, indent, before, line=SECTION_LINE_PT)
            add_spans(paragraph, (Span(title),), font=FONT, size=SECTION_PT)
        case Entry(spans):
            paragraph = _paragraph(cell, indent, _gap("entry", previous))
            add_spans(paragraph, entry_spans(spans), font=FONT, size=BODY_PT)
        case Meta(spans):
            paragraph = _paragraph(cell, indent, _gap("meta", previous))
            add_spans(paragraph, spans, font=FONT, size=BODY_PT, italic=True)
        case SkillLine(label, items):
            _add_skill(cell, label, items, indent=indent, sidebar=sidebar, previous=previous)
        case Bullet(spans):
            before = _gap(_bullet_kind(previous, sidebar=sidebar), previous)
            _add_bullet(cell, spans, indent=indent, before=before)
        case Paragraph(spans) if is_note(spans):
            inset = (indent[0] + NOTE_INDENT_IN, indent[1])
            paragraph = _paragraph(cell, inset, _gap("note", previous))
            add_spans(paragraph, spans, font=FONT, size=BODY_PT)
        case Paragraph(spans):
            paragraph = _paragraph(cell, indent, _gap("para", previous), line=PROSE_LINE_PT)
            add_spans(paragraph, spans, font=FONT, size=PROSE_PT)
        case _:  # pragma: no cover - mypy proves the block union is exhaustive
            assert_never(block)


def _add_bullet(
    cell: _Cell, spans: tuple[Span, ...], *, indent: tuple[float, float], before: float
) -> None:
    """One bulleted line: the glyph at the column edge, the text hanging past it."""
    paragraph = _paragraph(cell, indent, before, hanging=BULLET_HANGING_IN, style="List Bullet")
    add_spans(paragraph, spans, font=FONT, size=BODY_PT)


def _add_skill(  # noqa: PLR0913 - one skills line; each argument is a distinct input
    cell: _Cell,
    label: str,
    items: tuple[Span, ...],
    *,
    indent: tuple[float, float],
    sidebar: bool,
    previous: Block | None,
) -> None:
    """In the rail a skill group is a bold label over bulleted items; inline in the main column."""
    if not sidebar:
        paragraph = _paragraph(cell, indent, _gap("skill-inline", previous))
        add_spans(paragraph, (Span(f"{label}: "),), font=FONT, size=BODY_PT, bold=True)
        add_spans(paragraph, items, font=FONT, size=BODY_PT)
        return

    heading = _paragraph(cell, indent, _gap("skill-group", previous))
    add_spans(heading, (Span(label),), font=FONT, size=BODY_PT, bold=True)
    for index, item in enumerate(_skill_items(items)):
        kind = "skill-item" if index else "skill-item-first"
        _add_bullet(cell, item, indent=indent, before=GAP_PT[kind])


def _skill_items(items: tuple[Span, ...]) -> list[tuple[Span, ...]]:
    """Split ``a, b (c, d), e`` into items, ignoring commas inside parentheses."""
    entries: list[tuple[Span, ...]] = []
    current: list[Span] = []
    depth = 0
    buffer = ""
    style = Span("")

    def flush_buffer() -> None:
        nonlocal buffer
        if text := buffer.strip():
            current.append(Span(text, bold=style.bold, italic=style.italic))
        buffer = ""

    def flush_item() -> None:
        flush_buffer()
        if current:
            entries.append(tuple(current))
            current.clear()

    for span in items:
        flush_buffer()
        style = span
        for char in span.text:
            if char == "(":
                depth += 1
            elif char == ")":
                depth = max(depth - 1, 0)
            if char == "," and depth == 0:
                flush_item()
            else:
                buffer += char
    flush_item()
    return entries


# --- HTML -----------------------------------------------------------------------------------------
CSS = f"""
@page {{ size: Letter; margin: {PAGE_MARGIN_IN}in 0; }}
* {{ box-sizing: border-box; }}
/* Every gap is a margin-TOP, every bottom margin is 0, and line-height is absolute — the .docx
   sums adjacent spacing where CSS collapses it, and Word's line "multiple" is a ratio of the
   font's natural line box rather than of its size. Points mean the same thing to both. A word
   too long for its column (a long email or URL in the rail) breaks at the column's edge, as Word
   breaks it, instead of running across the rule into the other column. */
body {{ font-family: {FONT_STACK}; font-size: {BODY_PT}pt; line-height: {LINE_PT}pt;
        color: #111; margin: 0; overflow-wrap: anywhere; }}
/* Arial has two weights, regular and bold, which is also all a .docx run can say, so no rule
   here names a weight between them. An entry heading's bold comes from its spans, as in Word. */
h1, h2, h3, p, ul, li, div, aside, section {{ margin: 0; }}
/* The rule stops where the content does, exactly as the .docx table row does. */
.sheet {{ display: flex; align-items: stretch; }}
.rail {{ width: {SIDEBAR_WIDTH_IN}in; flex: 0 0 {SIDEBAR_WIDTH_IN}in; border-right: 1pt solid #000;
         padding: 0 {SIDEBAR_INDENT[1]}in 0 {SIDEBAR_INDENT[0]}in; }}
.main {{ width: {MAIN_WIDTH_IN}in; flex: 1 1 {MAIN_WIDTH_IN}in;
         padding: 0 {MAIN_INDENT[1]}in 0 {MAIN_INDENT[0]}in; }}
h1 {{ font-size: {NAME_PT}pt; font-weight: 400; line-height: {NAME_LINE_PT}pt;
      margin-top: {GAP_PT["name"]}pt; }}
h2 {{ font-size: {SECTION_PT}pt; font-weight: 400; line-height: {SECTION_LINE_PT}pt;
      margin-top: {GAP_PT["section"]}pt; }}
.rail h2 {{ margin-top: {GAP_PT["section-rail"]}pt; }}
h1 + h2 {{ margin-top: {GAP_PT["section-after-name"]}pt; }}
.subtitle + h2 {{ margin-top: {GAP_PT["section-after-subtitle"]}pt; }}
h3 {{ font-size: {BODY_PT}pt; font-weight: 400; margin-top: {GAP_PT["entry"]}pt; }}
.subtitle {{ font-size: {SUBTITLE_PT}pt; line-height: {SUBTITLE_LINE_PT}pt;
             margin-top: {GAP_PT["subtitle"]}pt; }}
.contact {{ font-size: {CONTACT_PT}pt; margin-top: {GAP_PT["contact"]}pt; }}
.meta {{ font-style: italic; margin-top: {GAP_PT["meta"]}pt; }}
.skill-group {{ font-weight: 700; margin-top: {GAP_PT["skill-group"]}pt; }}
.skill-inline {{ margin-top: {GAP_PT["skill-inline"]}pt; }}
p {{ font-size: {PROSE_PT}pt; line-height: {PROSE_LINE_PT}pt; margin-top: {GAP_PT["para"]}pt; }}
p.note {{ font-size: {BODY_PT}pt; line-height: {LINE_PT}pt; padding-left: {NOTE_INDENT_IN}in;
          margin-top: {GAP_PT["note"]}pt; }}
/* Every bullet draws its glyph at the column edge and hangs its text past it, with the left
   indent and negative first-line indent of the .docx "List Bullet" paragraphs. The glyph stays
   in the text flow: a positioned box is painted after everything else, so Chrome would write
   each bullet's text last and a parser would read every skill and accomplishment detached from
   its label or role. The first of a run stands off from what precedes it; the rest sit tighter. */
ul {{ padding: 0; list-style: none; }}
li {{ padding-left: {BULLET_HANGING_IN}in; text-indent: -{BULLET_HANGING_IN}in;
      margin-top: {GAP_PT["bullet"]}pt; }}
li:first-child {{ margin-top: {GAP_PT["bullet-first"]}pt; }}
li::before {{ content: "{BULLET_GLYPH}"; display: inline-block; width: {BULLET_HANGING_IN}in;
              text-indent: 0; }}
.rail li {{ margin-top: {GAP_PT["skill-item"]}pt; }}
.rail li:first-child {{ margin-top: {GAP_PT["skill-item-first"]}pt; }}
/* Both columns start at the page margin, as the first paragraph of each .docx cell does. */
.rail > :first-child, .main > :first-child {{ margin-top: 0; }}
"""


def render_html(
    document: Document, sidebar_sections: tuple[str, ...] = DEFAULT_SIDEBAR_SECTIONS
) -> str:
    """Render the polished layout as HTML, for print-to-PDF."""
    sidebar_groups, main_groups = split_columns(document, sidebar_sections)
    body = (
        '<div class="sheet">'
        f'<aside class="rail">{_html_groups(sidebar_groups, sidebar=True)}</aside>'
        f'<section class="main">{_html_groups(main_groups, sidebar=False)}</section>'
        "</div>"
    )
    return page(title_of(document), CSS, body)


def _html_groups(groups: list[SectionGroup], *, sidebar: bool) -> str:
    """Render the groups of one column, wrapping each run of bullets in a single list."""
    out: list[str] = []
    in_list = False
    for group in groups:
        if in_list:
            out.append("</ul>")
            in_list = False
        if group.title:
            out.append(f"<h2>{spans_to_html((Span(group.title),))}</h2>")
        for block in group.blocks:
            if isinstance(block, Bullet) != in_list:
                out.append("<ul>" if in_list is False else "</ul>")
                in_list = not in_list
            out.append(_html_block(block, sidebar=sidebar))
    if in_list:
        out.append("</ul>")
    return "".join(out)


def _html_block(block: GroupBlock, *, sidebar: bool) -> str:
    """Render one block; a bullet comes back as a bare ``<li>`` for the caller's list."""
    match block:
        case Name(spans):
            html = f"<h1>{spans_to_html(spans)}</h1>"
        case HeaderLine(spans):
            css = "contact" if sidebar else "subtitle"
            html = f'<div class="{css}">{spans_to_html(spans)}</div>'
        case Entry(spans):
            html = f"<h3>{spans_to_html(entry_spans(spans))}</h3>"
        case Meta(spans):
            html = f'<div class="meta">{spans_to_html(spans)}</div>'
        case SkillLine(label, items):
            html = _html_skill(label, items, sidebar=sidebar)
        case Bullet(spans):
            html = f"<li>{spans_to_html(spans)}</li>"
        case Paragraph(spans) if is_note(spans):
            html = f'<p class="note">{spans_to_html(spans)}</p>'
        case Paragraph(spans):
            html = f"<p>{spans_to_html(spans)}</p>"
        case _:  # pragma: no cover - mypy proves the block union is exhaustive
            assert_never(block)
    return html


def _html_skill(label: str, items: tuple[Span, ...], *, sidebar: bool) -> str:
    if not sidebar:
        return (
            f'<div class="skill-inline">{spans_to_html((Span(f"{label}: ", bold=True),))}'
            f"{spans_to_html(items)}</div>"
        )
    rows = "".join(f"<li>{spans_to_html(item)}</li>" for item in _skill_items(items))
    return f'<div class="skill-group">{spans_to_html((Span(label),))}</div><ul>{rows}</ul>'
